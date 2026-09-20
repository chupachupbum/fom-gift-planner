#!/usr/bin/env python3
"""
tests/test_challenger_cli_and_exporters.py

Empirical Challenger 2 Test Suite:
Validates CLI integration, Terminal presentation, and Excel multi-sheet export
for the Fields of Mistria Focus Recipes feature.
"""

import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import openpyxl

from fom_planner.cli import build_planner_parser, run_planner
from fom_planner.crafting import load_recipes
from fom_planner.data_loader import load_item_metadata, load_npc_preferences_from_json, load_recipe_sources
from fom_planner.exporters.excel_export import export_plan_to_excel
from fom_planner.exporters.terminal import print_terminal_plan
from fom_planner.models import SaveData
from fom_planner.optimizer import plan_daily_gift_bag, plan_max_relationship
from tests.test_gift_planner import create_synthetic_save_entries

DATA_RECIPES_PATH = REPO_ROOT / "data" / "recipes.json"
DATA_RECIPE_SOURCES_PATH = REPO_ROOT / "data" / "recipe_sources.json"


class TestCLIExecutionAndHelp(unittest.TestCase):
    """Empirical verification of CLI argument parsing and execution."""

    def test_cli_help_shows_recipe_sources(self):
        """python -m fom_planner.cli plan --help shows --recipe-sources."""
        parser = build_planner_parser()
        help_output = parser.format_help()
        self.assertIn("--recipe-sources", help_output)
        self.assertIn("recipe_sources.json", help_output)

    @patch("fom_planner.cli.find_latest_save", return_value=None)
    def test_cli_run_journal_strategy_with_recipe_sources(self, mock_save):
        """CLI execution: plan --strategy journal with --recipe-sources."""
        args = build_planner_parser().parse_args([
            "--strategy", "journal",
            "--recipe-sources", str(DATA_RECIPE_SOURCES_PATH),
            "--format", "terminal",
        ])
        buf = io.StringIO()
        with patch("sys.stdout", buf):
            run_planner(args)
        out = buf.getvalue()
        self.assertIn("FIELDS OF MISTRIA — DAILY GIFT BAG PLANNER", out)
        # Without save file, focus recipes should be silently omitted
        self.assertNotIn("🍳 FOCUS RECIPES", out)

    @patch("fom_planner.cli.find_latest_save", return_value=None)
    def test_cli_run_max_relationship_strategy_with_recipe_sources(self, mock_save):
        """CLI execution: plan --strategy max-relationship with --recipe-sources."""
        args = build_planner_parser().parse_args([
            "--strategy", "max-relationship",
            "--recipe-sources", str(DATA_RECIPE_SOURCES_PATH),
            "--format", "terminal",
        ])
        buf = io.StringIO()
        with patch("sys.stdout", buf):
            run_planner(args)
        out = buf.getvalue()
        self.assertIn("FIELDS OF MISTRIA — DAILY GIFT BAG PLANNER", out)
        self.assertNotIn("🍳 FOCUS RECIPES", out)

    def test_cli_nonexistent_recipe_sources_graceful_fallback(self):
        """CLI with nonexistent --recipe-sources runs without crashing."""
        args = build_planner_parser().parse_args([
            "--strategy", "journal",
            "--recipe-sources", "data/nonexistent_file_path.json",
            "--format", "terminal",
        ])
        buf = io.StringIO()
        with patch("sys.stdout", buf):
            run_planner(args)
        out = buf.getvalue()
        self.assertIn("FIELDS OF MISTRIA — DAILY GIFT BAG PLANNER", out)


class TestTerminalPresentationEmpirical(unittest.TestCase):
    """Empirical verification of Terminal presentation formatting and rules."""

    def setUp(self):
        self.recipes = load_recipes()
        self.cooking_keys = [k for k, v in self.recipes.items() if v.get("source") == "cooking"]
        self.sources = load_recipe_sources()
        self.raw = create_synthetic_save_entries(bag_items={"bread": 1})

    def _render_terminal(self, save, plan_results) -> str:
        buf = io.StringIO()
        with patch("sys.stdout", buf):
            print_terminal_plan(save, plan_results)
        return buf.getvalue()

    def test_terminal_without_save_file_silently_omitted(self):
        """Terminal output when run without save file must silently omit FOCUS RECIPES."""
        res = plan_daily_gift_bag(save=None, recipe_sources=self.sources)
        out = self._render_terminal(None, res)
        self.assertNotIn("🍳 FOCUS RECIPES", out)

    def test_terminal_some_recipes_unlocked_formatting_and_width(self):
        """Terminal presentation with 10 unlocked recipes displays header, table, and 86-char dividers."""
        pdict = json.loads(self.raw["player"])
        pdict["recipe_unlocks"] = self.cooking_keys[:10]
        self.raw["player"] = json.dumps(pdict)
        save = SaveData(Path("test_some.sav"), self.raw)

        res = plan_daily_gift_bag(save=save, recipe_sources=self.sources)
        out = self._render_terminal(save, res)

        self.assertIn("🍳 FOCUS RECIPES (10/163 cooking recipes unlocked — 6.1%):", out)
        self.assertIn("━" * 86, out)
        self.assertIn("═" * 86, out)

        # Check table headers and alignment
        expected_header = f" {'Rank':<5} │ {'Recipe Name':<24} │ {'Impact':<15} │ {'Unlock Source'}"
        self.assertIn(expected_header, out)

        # Verify dividers match 86 characters
        for line in out.splitlines():
            if line.startswith("━"):
                self.assertEqual(len(line), 86)
            elif line.startswith("═"):
                self.assertEqual(len(line), 86)

    def test_terminal_all_recipes_unlocked_celebration_banner(self):
        """Terminal presentation with all 163 recipes unlocked displays celebration banner."""
        pdict = json.loads(self.raw["player"])
        pdict["recipe_unlocks"] = self.cooking_keys
        self.raw["player"] = json.dumps(pdict)
        save = SaveData(Path("test_all.sav"), self.raw)

        res = plan_daily_gift_bag(save=save, recipe_sources=self.sources)
        out = self._render_terminal(save, res)

        self.assertIn("🍳 FOCUS RECIPES (163/163 cooking recipes unlocked — 100.0%):", out)
        self.assertIn("  🎉 All cooking recipes unlocked!", out)
        self.assertIn("━" * 86, out)

    def test_terminal_zero_cooking_recipes_unlocked(self):
        """Terminal presentation when save has unlocked items but 0 cooking recipes."""
        pdict = json.loads(self.raw["player"])
        # 'flour' is a milling recipe, not cooking
        pdict["recipe_unlocks"] = ["flour"]
        self.raw["player"] = json.dumps(pdict)
        save = SaveData(Path("test_0_cooking.sav"), self.raw)

        res = plan_daily_gift_bag(save=save, recipe_sources=self.sources)
        out = self._render_terminal(save, res)

        self.assertIn("🍳 FOCUS RECIPES (0/163 cooking recipes unlocked — 0.0%):", out)
        self.assertIn("━" * 86, out)
        self.assertIn("in 0 gifts", out)


class TestExcelExportEmpirical(unittest.TestCase):
    """Empirical inspection of Excel workbook structure using openpyxl."""

    def setUp(self):
        self.recipes = load_recipes()
        self.cooking_keys = [k for k, v in self.recipes.items() if v.get("source") == "cooking"]
        self.sources = load_recipe_sources()
        self.npcs, _ = load_npc_preferences_from_json(Path("data/item_data.json"))
        self.meta = load_item_metadata(Path("assets/fiddle"), Path("data/item_data.json"))

        raw = create_synthetic_save_entries(bag_items={"bread": 2, "apple": 5})
        pdict = json.loads(raw["player"])
        pdict["recipe_unlocks"] = self.cooking_keys[:10]
        raw["player"] = json.dumps(pdict)
        self.save = SaveData(Path("test_excel.sav"), raw)

    def test_excel_sheet_5_focus_recipes_structure(self):
        """Inspects that Sheet 5 exists, is named 'Focus Recipes', and matches expected headers and types."""
        plan = plan_daily_gift_bag(
            save=self.save,
            npc_gift_definitions=self.npcs,
            item_metadata=self.meta,
            recipe_sources=self.sources,
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_plan.xlsx"
            export_plan_to_excel(self.save, plan, self.npcs, self.meta, out_file)

            self.assertTrue(out_file.exists())
            wb = openpyxl.load_workbook(str(out_file))

            # 1. Total sheets must be 5
            self.assertEqual(len(wb.sheetnames), 5)

            # 2. Sheet 5 must be named "Focus Recipes"
            self.assertEqual(wb.sheetnames[4], "Focus Recipes")

            ws = wb["Focus Recipes"]

            # 3. Headers match ["Rank", "Recipe Name", "Impact", "Unlock Source"]
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(headers, ["Rank", "Recipe Name", "Impact", "Unlock Source"])

            # 4. Check data rows (at least 1 row, up to 10)
            data_rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertGreater(len(data_rows), 0)
            self.assertLessEqual(len(data_rows), 10)

            for idx, (rank, name, impact, source) in enumerate(data_rows, start=1):
                self.assertEqual(rank, idx)
                self.assertIsInstance(rank, int)
                self.assertIsInstance(name, str)
                self.assertGreater(len(name), 0)
                self.assertIsInstance(impact, int)
                self.assertGreaterEqual(impact, 0)
                self.assertIsInstance(source, str)
                self.assertGreater(len(source), 0)

            # 5. Styling and freeze panes
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertIsNotNone(ws.auto_filter.ref)

    def test_excel_sheet_5_max_relationship_strategy(self):
        """Inspects that Sheet 5 is generated properly with max-relationship strategy."""
        plan = plan_max_relationship(
            save=self.save,
            npc_gift_definitions=self.npcs,
            item_metadata=self.meta,
            recipe_sources=self.sources,
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_plan_max_rel.xlsx"
            export_plan_to_excel(self.save, plan, self.npcs, self.meta, out_file)

            wb = openpyxl.load_workbook(str(out_file))
            self.assertIn("Focus Recipes", wb.sheetnames)
            ws = wb["Focus Recipes"]
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(headers, ["Rank", "Recipe Name", "Impact", "Unlock Source"])


if __name__ == "__main__":
    unittest.main()
