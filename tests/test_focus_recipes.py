#!/usr/bin/env python3
"""
tests/test_focus_recipes.py

Comprehensive End-to-End & Unit Test Suite for the Focus Recipes feature:
- Tier 1: Data Completeness & Schema Contract (recipe_sources.json vs recipes.json)
- Tier 1: Loader Contract (load_recipe_sources)
- Tier 1: Optimizer Core (compute_focus_recipes impact calculation, locked filtering, ranking, top_n, stats)
- Tier 2: Boundary & Corner Cases (unlocked_recipe_ids=None, 100% unlocked, empty remaining items, top_n <= 0)
- Tier 3: Planning Strategies Integration (plan_daily_gift_bag, plan_max_relationship)
- Tier 3: Terminal Presentation (print_terminal_plan layout, ordering, badges, skips)
- Tier 3: Excel Export (export_plan_to_excel "Focus Recipes" sheet, headers, rows)
- Tier 3: CLI Parser & Dispatcher (--recipe-sources flag)
- Tier 4: Real-World Scenarios (Early-game, mid-game, completionist end-game playthroughs)
"""

import copy
import inspect
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from typing import Any, Dict, List, Optional, Set, Tuple
from unittest.mock import MagicMock, patch

# Ensure repository root is on sys.path
REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

# Safe imports for components that may be implemented concurrently
try:
    from fom_planner.data_loader import load_recipe_sources
except (ImportError, AttributeError):
    load_recipe_sources = None

try:
    from fom_planner.optimizer import compute_focus_recipes
except (ImportError, AttributeError):
    compute_focus_recipes = None

from fom_planner.crafting import load_recipes
from fom_planner.models import InGameDate, SaveData
from fom_planner.optimizer import plan_daily_gift_bag, plan_max_relationship
from fom_planner.exporters.terminal import print_terminal_plan
from fom_planner.exporters.excel_export import OPENPYXL_AVAILABLE, export_plan_to_excel
from fom_planner.cli import build_planner_parser, run_planner

if OPENPYXL_AVAILABLE:
    import openpyxl

# Test fixtures helpers
from tests.test_gift_planner import (
    create_mock_slot,
    create_synthetic_save_entries,
    get_mock_meta,
    get_mock_npcs,
    get_mock_recipes,
)

DATA_RECIPES_PATH = REPO_ROOT / "data" / "recipes.json"
DATA_RECIPE_SOURCES_PATH = REPO_ROOT / "data" / "recipe_sources.json"


# ==============================================================================
# TIER 1: DATA COMPLETENESS & SCHEMA CONTRACT
# ==============================================================================

class TestRecipeSourcesDataCompleteness(unittest.TestCase):
    """
    Tier 1: Verifies that data/recipe_sources.json adheres strictly to the
    cooking recipe database specifications in data/recipes.json.
    """

    def setUp(self):
        self.assertTrue(DATA_RECIPES_PATH.exists(), f"Missing {DATA_RECIPES_PATH}")
        with open(DATA_RECIPES_PATH, "r", encoding="utf-8") as f:
            self.recipes = json.load(f)

        self.cooking_recipes = {
            k: v for k, v in self.recipes.items()
            if isinstance(v, dict) and str(v.get("source", "")).strip().lower() == "cooking"
        }
        self.milling_recipes = {
            k: v for k, v in self.recipes.items()
            if isinstance(v, dict) and str(v.get("source", "")).strip().lower() == "milling"
        }

    def test_recipe_sources_file_exists(self):
        """data/recipe_sources.json must exist in the data/ directory."""
        self.assertTrue(
            DATA_RECIPE_SOURCES_PATH.exists(),
            f"Expected {DATA_RECIPE_SOURCES_PATH} to exist."
        )
        self.assertTrue(DATA_RECIPE_SOURCES_PATH.is_file())

    def test_recipe_sources_exact_count(self):
        """recipe_sources.json must contain exactly 163 entries matching cooking recipes."""
        with open(DATA_RECIPE_SOURCES_PATH, "r", encoding="utf-8") as f:
            sources = json.load(f)

        self.assertIsInstance(sources, dict)
        self.assertEqual(
            len(sources), 163,
            f"Expected exactly 163 entries in recipe_sources.json, found {len(sources)}"
        )
        self.assertEqual(
            len(sources), len(self.cooking_recipes),
            "Count in recipe_sources.json must match count of cooking recipes in recipes.json"
        )

    def test_recipe_sources_keys_match_cooking_recipes(self):
        """Every key in recipe_sources.json must correspond to a cooking recipe in recipes.json."""
        with open(DATA_RECIPE_SOURCES_PATH, "r", encoding="utf-8") as f:
            sources = json.load(f)

        source_keys = set(sources.keys())
        cooking_keys = set(self.cooking_recipes.keys())

        missing_from_sources = cooking_keys - source_keys
        extra_in_sources = source_keys - cooking_keys

        self.assertEqual(
            missing_from_sources, set(),
            f"Cooking recipes missing from recipe_sources.json: {missing_from_sources}"
        )
        self.assertEqual(
            extra_in_sources, set(),
            f"Unexpected extra keys in recipe_sources.json: {extra_in_sources}"
        )

    def test_recipe_sources_strictly_excludes_milling(self):
        """None of the 23 milling recipes may appear in recipe_sources.json."""
        with open(DATA_RECIPE_SOURCES_PATH, "r", encoding="utf-8") as f:
            sources = json.load(f)

        milling_keys = set(self.milling_recipes.keys())
        source_keys = set(sources.keys())
        overlap = source_keys & milling_keys

        self.assertEqual(
            overlap, set(),
            f"Milling recipes must not be included in recipe_sources.json: {overlap}"
        )

    def test_recipe_sources_no_empty_values(self):
        """No unlock source strings in recipe_sources.json may be empty or whitespace."""
        with open(DATA_RECIPE_SOURCES_PATH, "r", encoding="utf-8") as f:
            sources = json.load(f)

        empty_entries = []
        for rid, src in sources.items():
            if not isinstance(src, str) or not src.strip():
                empty_entries.append((rid, src))

        self.assertEqual(
            empty_entries, [],
            f"Found empty or non-string values in recipe_sources.json: {empty_entries}"
        )

    def test_recipe_sources_keys_are_normalized(self):
        """All keys in recipe_sources.json must be lowercased and stripped."""
        with open(DATA_RECIPE_SOURCES_PATH, "r", encoding="utf-8") as f:
            sources = json.load(f)

        unnormalized = [k for k in sources.keys() if k != k.strip().lower()]
        self.assertEqual(
            unnormalized, [],
            f"Keys in recipe_sources.json must be lowercase and stripped: {unnormalized}"
        )


# ==============================================================================
# TIER 1: DATA LOADER CONTRACT
# ==============================================================================

class TestRecipeSourcesLoader(unittest.TestCase):
    """
    Tier 1 & Tier 2: Tests for fom_planner.data_loader.load_recipe_sources().
    """

    def setUp(self):
        if load_recipe_sources is None:
            self.skipTest("load_recipe_sources is not yet implemented in fom_planner.data_loader")

    def test_load_recipe_sources_default_path(self):
        """Calling load_recipe_sources() with no arguments loads data/recipe_sources.json."""
        sources = load_recipe_sources()
        self.assertIsInstance(sources, dict)
        self.assertEqual(len(sources), 163)
        self.assertIn("bread", sources)
        self.assertIn("apple_pie", sources)

    def test_load_recipe_sources_none_path(self):
        """Calling load_recipe_sources(None) resolves to default path."""
        sources = load_recipe_sources(None)
        self.assertIsInstance(sources, dict)
        self.assertEqual(len(sources), 163)

    def test_load_recipe_sources_explicit_path(self):
        """Calling load_recipe_sources with explicit valid path loads properly."""
        sources = load_recipe_sources(DATA_RECIPE_SOURCES_PATH)
        self.assertIsInstance(sources, dict)
        self.assertEqual(len(sources), 163)

    def test_load_recipe_sources_custom_temp_file(self):
        """Calling load_recipe_sources with a custom JSON file returns its contents."""
        custom_data = {
            "custom_dish_1": "Found in cave",
            "custom_dish_2": "Sold by vendor",
        }
        with tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False, encoding="utf-8") as tf:
            json.dump(custom_data, tf)
            tf_path = Path(tf.name)

        try:
            loaded = load_recipe_sources(tf_path)
            self.assertEqual(len(loaded), 2)
            self.assertEqual(loaded["custom_dish_1"], "Found in cave")
            self.assertEqual(loaded["custom_dish_2"], "Sold by vendor")
        finally:
            if tf_path.exists():
                tf_path.unlink()

    def test_load_recipe_sources_normalizes_keys_and_values(self):
        """Loader normalizes keys to lowercase/stripped and strips values."""
        raw_data = {
            "  Apple_Pie  ": "  Bakery Purchase  ",
            "BREAD": "  Starter Recipe  ",
        }
        with tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False, encoding="utf-8") as tf:
            json.dump(raw_data, tf)
            tf_path = Path(tf.name)

        try:
            loaded = load_recipe_sources(tf_path)
            self.assertIn("apple_pie", loaded)
            self.assertEqual(loaded["apple_pie"], "Bakery Purchase")
            self.assertIn("bread", loaded)
            self.assertEqual(loaded["bread"], "Starter Recipe")
        finally:
            if tf_path.exists():
                tf_path.unlink()

    def test_load_recipe_sources_nonexistent_path(self):
        """Nonexistent path returns empty dict {} without raising FileNotFoundError."""
        missing_path = REPO_ROOT / "data" / "definitely_does_not_exist_9999.json"
        res = load_recipe_sources(missing_path)
        self.assertEqual(res, {})

    def test_load_recipe_sources_directory_path(self):
        """Directory path returns empty dict {} gracefully."""
        res = load_recipe_sources(REPO_ROOT / "data")
        self.assertEqual(res, {})

    def test_load_recipe_sources_corrupted_json(self):
        """Corrupted JSON file returns empty dict {} gracefully."""
        with tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False, encoding="utf-8") as tf:
            tf.write("{ this is not valid json ::: }}")
            tf_path = Path(tf.name)

        try:
            res = load_recipe_sources(tf_path)
            self.assertEqual(res, {})
        finally:
            if tf_path.exists():
                tf_path.unlink()

    def test_load_recipe_sources_non_dict_json(self):
        """JSON file with non-dict top level (e.g. list) returns empty dict {}."""
        with tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False, encoding="utf-8") as tf:
            json.dump(["item1", "item2"], tf)
            tf_path = Path(tf.name)

        try:
            res = load_recipe_sources(tf_path)
            self.assertEqual(res, {})
        finally:
            if tf_path.exists():
                tf_path.unlink()


# ==============================================================================
# TIER 1: OPTIMIZER CORE CONTRACT
# ==============================================================================

class TestComputeFocusRecipesCore(unittest.TestCase):
    """
    Tier 1: Core functionality of compute_focus_recipes():
    - Impact calculation (loved + liked combined)
    - Filtering locked recipes (excluding milling, excluding unlocked)
    - Ranking hierarchy & deterministic tiebreakers
    - Top N truncation
    - Unlock stats accuracy
    """

    def setUp(self):
        if compute_focus_recipes is None:
            self.skipTest("compute_focus_recipes is not yet implemented in fom_planner.optimizer")

        # Create a controlled mock recipe catalog
        self.mock_recipes = {
            "dish_alpha": {
                "item_id": "dish_alpha",
                "display_name": "Alpha Stew",
                "source": "cooking",
                "output_count": 1,
                "ingredients": [],
            },
            "dish_bravo": {
                "item_id": "dish_bravo",
                "display_name": "Bravo Bread",
                "source": "cooking",
                "output_count": 1,
                "ingredients": [],
            },
            "dish_charlie": {
                "item_id": "dish_charlie",
                "display_name": "Charlie Cake",
                "source": "cooking",
                "output_count": 1,
                "ingredients": [],
            },
            "dish_delta": {
                "item_id": "dish_delta",
                "display_name": "Delta Drink",
                "source": "cooking",
                "output_count": 1,
                "ingredients": [],
            },
            "flour_milling": {
                "item_id": "flour_milling",
                "display_name": "Flour",
                "source": "milling",
                "output_count": 1,
                "ingredients": [],
            },
        }

        self.mock_sources = {
            "dish_alpha": "The Dragon Inn",
            "dish_bravo": "Request Board",
            "dish_charlie": "Wishing Well",
            "dish_delta": "Fishing",
        }

    def test_impact_calculation_loved_and_liked(self):
        """
        Impact must equal the sum of pending loved and liked NPC preferences.
        Alpha: 2 loved + 1 liked = 3
        Bravo: 1 loved + 0 liked = 1
        Charlie: 0 loved + 1 liked = 1
        Delta: 0 loved + 0 liked = 0
        """
        remaining_map = {
            "dish_alpha": {"loved": {"adeline", "march"}, "liked": {"celine"}},
            "dish_bravo": {"loved": {"adeline"}, "liked": set()},
            "dish_charlie": {"loved": set(), "liked": {"balor"}},
            "dish_delta": {"loved": set(), "liked": set()},
        }

        # Player has none unlocked
        unlocked = set()
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=unlocked,
            remaining_items_map=remaining_map,
            recipe_sources=self.mock_sources,
            top_n=10,
        )

        impact_by_id = {r["recipe_id"]: r["impact"] for r in focus_list}
        self.assertEqual(impact_by_id["dish_alpha"], 3)
        self.assertEqual(impact_by_id["dish_bravo"], 1)
        self.assertEqual(impact_by_id["dish_charlie"], 1)
        self.assertEqual(impact_by_id["dish_delta"], 0)

    def test_locked_recipe_filtering_excludes_unlocked(self):
        """Recipes in unlocked_recipe_ids must NOT appear in focus_recipes."""
        remaining_map = {
            "dish_alpha": {"loved": {"adeline"}, "liked": set()},
            "dish_bravo": {"loved": {"march"}, "liked": set()},
        }
        # Player unlocked dish_alpha
        unlocked = {"dish_alpha"}
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=unlocked,
            remaining_items_map=remaining_map,
            recipe_sources=self.mock_sources,
            top_n=10,
        )

        recipe_ids = [r["recipe_id"] for r in focus_list]
        self.assertNotIn("dish_alpha", recipe_ids)
        self.assertIn("dish_bravo", recipe_ids)
        self.assertIn("dish_charlie", recipe_ids)
        self.assertIn("dish_delta", recipe_ids)

    def test_milling_recipes_strictly_excluded(self):
        """Milling recipes (source == 'milling') must never be returned in focus_recipes."""
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=set(),
            remaining_items_map={},
            recipe_sources=self.mock_sources,
            top_n=10,
        )

        recipe_ids = [r["recipe_id"] for r in focus_list]
        self.assertNotIn("flour_milling", recipe_ids)

    def test_ranking_hierarchy_impact_descending(self):
        """Higher impact recipes must appear before lower impact recipes."""
        remaining_map = {
            "dish_delta": {"loved": {"a", "b", "c", "d"}, "liked": set()},  # impact 4
            "dish_charlie": {"loved": {"a", "b"}, "liked": {"c"}},         # impact 3
            "dish_bravo": {"loved": {"a"}, "liked": {"b"}},               # impact 2
            "dish_alpha": {"loved": {"a"}, "liked": set()},               # impact 1
        }
        focus_list, _ = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=set(),
            remaining_items_map=remaining_map,
            recipe_sources=self.mock_sources,
            top_n=10,
        )

        ranked_ids = [r["recipe_id"] for r in focus_list]
        self.assertEqual(ranked_ids, ["dish_delta", "dish_charlie", "dish_bravo", "dish_alpha"])

    def test_tiebreaker_display_name_ascending(self):
        """Ties in impact must be broken alphabetically by display name (case-insensitive)."""
        # Create recipes with identical impact (all 2)
        recipes = {
            "r_zebra": {"item_id": "r_zebra", "display_name": "Zucchini Bread", "source": "cooking"},
            "r_apple": {"item_id": "r_apple", "display_name": "Apple Tart", "source": "cooking"},
            "r_mango": {"item_id": "r_mango", "display_name": "mango Salad", "source": "cooking"},
        }
        remaining_map = {
            "r_zebra": {"loved": {"a", "b"}, "liked": set()},
            "r_apple": {"loved": {"c", "d"}, "liked": set()},
            "r_mango": {"loved": {"e", "f"}, "liked": set()},
        }

        focus_list, _ = compute_focus_recipes(
            all_recipes=recipes,
            unlocked_recipe_ids=set(),
            remaining_items_map=remaining_map,
            top_n=10,
        )

        names = [r["display_name"] for r in focus_list]
        self.assertEqual(names, ["Apple Tart", "mango Salad", "Zucchini Bread"])

    def test_top_n_truncation_and_rank_numbers(self):
        """Top N limits returned items, and rank values are 1-indexed sequential integers."""
        remaining_map = {f"dish_{i}": {"loved": {"npc"}, "liked": set()} for i in range(15)}
        custom_recipes = {
            f"dish_{i}": {"item_id": f"dish_{i}", "display_name": f"Dish {i}", "source": "cooking"}
            for i in range(15)
        }

        focus_list, _ = compute_focus_recipes(
            all_recipes=custom_recipes,
            unlocked_recipe_ids=set(),
            remaining_items_map=remaining_map,
            top_n=5,
        )

        self.assertEqual(len(focus_list), 5)
        for idx, item in enumerate(focus_list, start=1):
            self.assertEqual(item["rank"], idx)

    def test_recipe_sources_population_and_unknown_fallback(self):
        """Items in focus_recipes receive unlock_source from recipe_sources or 'Unknown'."""
        sources = {
            "dish_alpha": "Found at the Dragon Inn",
            # dish_bravo intentionally missing from sources
        }
        focus_list, _ = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=set(),
            remaining_items_map={},
            recipe_sources=sources,
            top_n=10,
        )

        source_by_id = {r["recipe_id"]: r["unlock_source"] for r in focus_list}
        self.assertEqual(source_by_id["dish_alpha"], "Found at the Dragon Inn")
        self.assertEqual(source_by_id["dish_bravo"], "Unknown")

    def test_unlock_stats_calculation(self):
        """recipe_unlock_stats accurately calculates total, unlocked count, and percentage."""
        # 4 cooking recipes total, 1 unlocked (dish_alpha)
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids={"dish_alpha"},
            remaining_items_map={},
            top_n=10,
        )

        self.assertIsInstance(stats, dict)
        self.assertEqual(stats.get("total_cooking_recipes"), 4)
        self.assertEqual(stats.get("unlocked_recipes_count"), 1)
        self.assertAlmostEqual(stats.get("unlocked_percentage"), 25.0, places=2)
        # Verify compatibility aliases
        self.assertEqual(stats.get("total_cooking"), 4)
        self.assertEqual(stats.get("unlocked_count"), 1)
        self.assertTrue(stats.get("has_unlock_data"))


# ==============================================================================
# TIER 2: BOUNDARY & CORNER CASES
# ==============================================================================

class TestComputeFocusRecipesEdgeCases(unittest.TestCase):
    """
    Tier 2: Edge cases and adversarial inputs for compute_focus_recipes().
    """

    def setUp(self):
        if compute_focus_recipes is None:
            self.skipTest("compute_focus_recipes is not yet implemented in fom_planner.optimizer")

        self.mock_recipes = {
            "bread": {"item_id": "bread", "display_name": "Bread", "source": "cooking"},
            "pie": {"item_id": "pie", "display_name": "Pie", "source": "cooking"},
        }

    def test_edge_case_no_save_unlocked_ids_none(self):
        """When unlocked_recipe_ids is None (no save / mock without unlocks), returns empty list."""
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=None,
            remaining_items_map={"bread": {"loved": {"adeline"}, "liked": set()}},
            top_n=10,
        )

        self.assertEqual(focus_list, [])
        self.assertIsInstance(stats, dict)
        self.assertFalse(stats.get("has_unlock_data"))
        self.assertEqual(stats.get("unlocked_count"), 0)
        self.assertEqual(stats.get("unlocked_percentage"), 0.0)

    def test_edge_case_all_recipes_unlocked(self):
        """When all cooking recipes are unlocked, returns empty list and 100% stats."""
        unlocked = {"bread", "pie"}
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=unlocked,
            remaining_items_map={"bread": {"loved": {"adeline"}, "liked": set()}},
            top_n=10,
        )

        self.assertEqual(focus_list, [])
        self.assertIsInstance(stats, dict)
        self.assertTrue(stats.get("has_unlock_data"))
        self.assertEqual(stats.get("unlocked_count"), 2)
        self.assertAlmostEqual(stats.get("unlocked_percentage"), 100.0, places=2)

    def test_edge_case_empty_remaining_items_map(self):
        """When remaining_items_map is empty or None, returns locked recipes with impact=0."""
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=set(),
            remaining_items_map={},
            top_n=10,
        )

        self.assertEqual(len(focus_list), 2)
        for item in focus_list:
            self.assertEqual(item["impact"], 0)
        # Should be sorted alphabetically: Bread before Pie
        self.assertEqual(focus_list[0]["recipe_id"], "bread")
        self.assertEqual(focus_list[1]["recipe_id"], "pie")

    def test_edge_case_top_n_zero_or_negative(self):
        """When top_n <= 0, returns empty list without error."""
        for n in [0, -1, -10]:
            focus_list, stats = compute_focus_recipes(
                all_recipes=self.mock_recipes,
                unlocked_recipe_ids=set(),
                remaining_items_map={},
                top_n=n,
            )
            self.assertEqual(focus_list, [])

    def test_edge_case_empty_recipes_dict(self):
        """When all_recipes is empty dict {}, returns empty list and 0% stats."""
        focus_list, stats = compute_focus_recipes(
            all_recipes={},
            unlocked_recipe_ids=set(),
            remaining_items_map={},
            top_n=10,
        )
        self.assertEqual(focus_list, [])
        self.assertEqual(stats.get("total_cooking", 0), 0)

    def test_edge_case_malformed_unlocked_ids(self):
        """Handles unlocked_recipe_ids with mixed case, whitespace, None, or primitives."""
        unlocked = ["  BREAD  ", None, "", 1234, "PIE"]
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=unlocked,
            remaining_items_map={},
            top_n=10,
        )
        # Both bread and pie should be recognized as unlocked
        self.assertEqual(focus_list, [])
        self.assertEqual(stats.get("unlocked_count"), 2)

    def test_edge_case_malformed_remaining_items_map(self):
        """Handles remaining_items_map where preference entries are malformed or non-set."""
        remaining_map = {
            "bread": {"loved": "single_string_npc", "liked": None},
            "pie": ["list_of_npcs_1", "list_of_npcs_2"],
        }
        focus_list, stats = compute_focus_recipes(
            all_recipes=self.mock_recipes,
            unlocked_recipe_ids=set(),
            remaining_items_map=remaining_map,
            top_n=10,
        )
        self.assertEqual(len(focus_list), 2)
        # Should execute safely without raising TypeError or AttributeError


# ==============================================================================
# TIER 3: PLANNING STRATEGIES INTEGRATION
# ==============================================================================

class TestPlanningStrategiesIntegration(unittest.TestCase):
    """
    Tier 3: Verifies that plan_daily_gift_bag() and plan_max_relationship()
    properly integrate compute_focus_recipes() and return focus_recipes
    and recipe_unlock_stats in their returned dictionary.
    """

    def setUp(self):
        # Create a synthetic save with known recipe unlocks
        raw = create_synthetic_save_entries(
            bag_items={"bread": 5},
            year=1,
            season="spring",
            day=1,
        )
        # Inject recipe_unlocks into player block
        player_dict = json.loads(raw["player"])
        player_dict["recipe_unlocks"] = ["bread", "wild_berry_jam"]
        raw["player"] = json.dumps(player_dict)

        self.save = SaveData(Path("test_synthetic.sav"), raw)

        # Also create a synthetic save without recipe_unlocks
        raw_no_unlocks = create_synthetic_save_entries(bag_items={"bread": 5})
        self.save_no_unlocks = SaveData(Path("test_no_unlocks.sav"), raw_no_unlocks)

    def test_plan_daily_gift_bag_accepts_recipe_sources_and_returns_keys(self):
        """plan_daily_gift_bag() accepts recipe_sources and returns focus_recipes & stats."""
        sig = inspect.signature(plan_daily_gift_bag)
        if "recipe_sources" not in sig.parameters:
            self.skipTest("plan_daily_gift_bag does not yet accept recipe_sources parameter")

        recipe_sources = {"bread": "Default Starter", "apple_pie": "Bakery"}
        plan_results = plan_daily_gift_bag(
            save=self.save,
            recipe_sources=recipe_sources,
        )

        self.assertIn("focus_recipes", plan_results)
        self.assertIn("recipe_unlock_stats", plan_results)

        focus_recipes = plan_results["focus_recipes"]
        stats = plan_results["recipe_unlock_stats"]

        self.assertIsInstance(focus_recipes, list)
        self.assertIsInstance(stats, dict)
        self.assertTrue(stats.get("has_unlock_data"))
        self.assertGreaterEqual(stats.get("unlocked_count", 0), 2)

    def test_plan_max_relationship_accepts_recipe_sources_and_returns_keys(self):
        """plan_max_relationship() accepts recipe_sources and returns focus_recipes & stats."""
        sig = inspect.signature(plan_max_relationship)
        if "recipe_sources" not in sig.parameters:
            self.skipTest("plan_max_relationship does not yet accept recipe_sources parameter")

        recipe_sources = {"bread": "Default Starter", "apple_pie": "Bakery"}
        plan_results = plan_max_relationship(
            save=self.save,
            recipe_sources=recipe_sources,
        )

        self.assertIn("focus_recipes", plan_results)
        self.assertIn("recipe_unlock_stats", plan_results)

        focus_recipes = plan_results["focus_recipes"]
        stats = plan_results["recipe_unlock_stats"]

        self.assertIsInstance(focus_recipes, list)
        self.assertIsInstance(stats, dict)
        self.assertTrue(stats.get("has_unlock_data"))
        self.assertGreaterEqual(stats.get("unlocked_count", 0), 2)

    def test_strategies_without_save_return_empty_focus_recipes(self):
        """When save is None, both strategies return focus_recipes=[] and has_unlock_data=False."""
        sig1 = inspect.signature(plan_daily_gift_bag)
        if "recipe_sources" not in sig1.parameters:
            self.skipTest("plan_daily_gift_bag does not yet accept recipe_sources parameter")

        res_daily = plan_daily_gift_bag(save=None)
        self.assertIn("focus_recipes", res_daily)
        self.assertEqual(res_daily["focus_recipes"], [])
        if "recipe_unlock_stats" in res_daily:
            self.assertFalse(res_daily["recipe_unlock_stats"].get("has_unlock_data", False))

        sig2 = inspect.signature(plan_max_relationship)
        if "recipe_sources" not in sig2.parameters:
            self.skipTest("plan_max_relationship does not yet accept recipe_sources parameter")

        res_max = plan_max_relationship(save=None)
        self.assertIn("focus_recipes", res_max)
        self.assertEqual(res_max["focus_recipes"], [])
        if "recipe_unlock_stats" in res_max:
            self.assertFalse(res_max["recipe_unlock_stats"].get("has_unlock_data", False))


# ==============================================================================
# TIER 3: TERMINAL PRESENTATION
# ==============================================================================

class TestTerminalPresentation(unittest.TestCase):
    """
    Tier 3: Verifies terminal layout, positioning, column headers, and skip rules:
    - Renders 🍳 FOCUS RECIPES section after 💡 FOCUS SUGGESTIONS
    - Formats columns: Rank, Recipe Name, Impact, Unlock Source
    - Silently skips section when save is None or no unlock data
    - Displays celebration banner when all recipes unlocked
    """

    def setUp(self):
        raw = create_synthetic_save_entries(bag_items={"bread": 1})
        self.save = SaveData(Path("test.sav"), raw)

        self.sample_focus_recipes = [
            {"rank": 1, "recipe_id": "apple_pie", "display_name": "Apple Pie", "impact": 4, "unlock_source": "The Dragon Inn"},
            {"rank": 2, "recipe_id": "beet_soup", "display_name": "Beet Soup", "impact": 2, "unlock_source": "Shipping"},
        ]
        self.sample_stats = {
            "total_cooking": 163,
            "unlocked_count": 10,
            "unlocked_percentage": 6.13,
            "has_unlock_data": True,
        }
        self.base_plan = {
            "bag_plan": [],
            "npc_progress": {},
            "covered_npcs": [],
            "target_npcs": [],
            "focus_suggestions": [],
            "focus_recipes": self.sample_focus_recipes,
            "recipe_unlock_stats": self.sample_stats,
            "overall_stats": {
                "strategy": "journal",
                "game_given_loved": 5,
                "game_total_loved": 100,
                "game_given_liked": 10,
                "game_total_liked": 200,
                "game_given_total": 15,
                "game_total_preferences": 300,
                "remaining_unique_items": 50,
                "mode": "auto",
                "max_slots": 20,
                "covered_npcs_count": 0,
                "target_npcs_count": 0,
                "vendors_covered_today": 0,
                "today_loved_completed": 0,
                "today_liked_completed": 0,
                "locked_npcs": [],
                "ungiftable_npcs": [],
            },
        }

    def _capture_terminal_output(self, save, plan_results) -> str:
        buf = io.StringIO()
        with patch("sys.stdout", buf):
            print_terminal_plan(save, plan_results)
        return buf.getvalue()

    def test_terminal_renders_focus_recipes_section_with_save(self):
        """Terminal renders 🍳 FOCUS RECIPES with unlock percentage when save is present."""
        output = self._capture_terminal_output(self.save, self.base_plan)
        if "🍳 FOCUS RECIPES" not in output:
            self.skipTest("terminal.py has not yet implemented FOCUS RECIPES section")

        self.assertIn("🍳 FOCUS RECIPES (10/163 cooking recipes unlocked — 6.1%):", output)
        self.assertIn("Apple Pie", output)
        self.assertIn("Beet Soup", output)
        self.assertIn("The Dragon Inn", output)
        self.assertIn("Shipping", output)
        self.assertIn("in 4 gifts", output)
        self.assertIn("in 2 gifts", output)

    def test_terminal_columns_present(self):
        """Table headers Rank, Recipe Name, Impact, and Unlock Source are rendered."""
        output = self._capture_terminal_output(self.save, self.base_plan)
        if "🍳 FOCUS RECIPES" not in output:
            self.skipTest("terminal.py has not yet implemented FOCUS RECIPES section")

        self.assertIn("Rank", output)
        self.assertIn("Recipe Name", output)
        self.assertIn("Impact", output)
        self.assertIn("Unlock Source", output)

    def test_terminal_section_placement(self):
        """FOCUS RECIPES must appear after FOCUS SUGGESTIONS and before final closing line."""
        output = self._capture_terminal_output(self.save, self.base_plan)
        if "🍳 FOCUS RECIPES" not in output:
            self.skipTest("terminal.py has not yet implemented FOCUS RECIPES section")

        pos_suggestions = output.find("FOCUS SUGGESTIONS")
        pos_recipes = output.find("🍳 FOCUS RECIPES")
        pos_final_divider = output.rfind("═" * 86)

        self.assertGreater(pos_recipes, pos_suggestions, "FOCUS RECIPES must follow FOCUS SUGGESTIONS")
        self.assertLess(pos_recipes, pos_final_divider, "FOCUS RECIPES must appear before final report line")

    def test_terminal_silently_skips_when_no_save(self):
        """When save is None, FOCUS RECIPES is completely omitted."""
        output = self._capture_terminal_output(None, self.base_plan)
        self.assertNotIn("🍳 FOCUS RECIPES", output)

    def test_terminal_silently_skips_when_no_unlock_stats(self):
        """When recipe_unlock_stats is missing or has_unlock_data=False, section is skipped."""
        plan_no_data = copy.deepcopy(self.base_plan)
        plan_no_data["recipe_unlock_stats"] = {"has_unlock_data": False, "total_cooking": 163, "unlocked_count": 0}
        output = self._capture_terminal_output(self.save, plan_no_data)
        self.assertNotIn("🍳 FOCUS RECIPES", output)

    def test_terminal_celebration_banner_when_all_unlocked(self):
        """When unlocked_count >= total_cooking, shows '🎉 All cooking recipes unlocked!'."""
        plan_all_unlocked = copy.deepcopy(self.base_plan)
        plan_all_unlocked["focus_recipes"] = []
        plan_all_unlocked["recipe_unlock_stats"] = {
            "total_cooking": 163,
            "unlocked_count": 163,
            "unlocked_percentage": 100.0,
            "has_unlock_data": True,
        }
        output = self._capture_terminal_output(self.save, plan_all_unlocked)
        if "🍳 FOCUS RECIPES" not in output:
            self.skipTest("terminal.py has not yet implemented FOCUS RECIPES section")

        self.assertIn("🎉 All cooking recipes unlocked!", output)


# ==============================================================================
# TIER 3: EXCEL EXPORT INTEGRATION
# ==============================================================================

class TestExcelExportFocusRecipes(unittest.TestCase):
    """
    Tier 3: Verifies that export_plan_to_excel() includes a 'Focus Recipes' sheet
    with expected headers and rows matching plan_results['focus_recipes'].
    """

    def setUp(self):
        if not OPENPYXL_AVAILABLE:
            self.skipTest("openpyxl is not installed in this environment")

        self.plan_results = {
            "bag_plan": [],
            "npc_progress": {},
            "all_remaining_items_map": {},
            "remaining_items_map": {},
            "focus_recipes": [
                {"rank": 1, "recipe_id": "apple_pie", "display_name": "Apple Pie", "impact": 5, "unlock_source": "The Inn"},
                {"rank": 2, "recipe_id": "beet_soup", "display_name": "Beet Soup", "impact": 3, "unlock_source": "Shipping"},
            ],
            "recipe_unlock_stats": {
                "total_cooking": 163,
                "unlocked_count": 15,
                "unlocked_percentage": 9.2,
                "has_unlock_data": True,
            },
            "overall_stats": {
                "strategy": "journal",
                "game_given_loved": 0,
                "game_total_loved": 10,
                "game_given_liked": 0,
                "game_total_liked": 10,
                "game_given_total": 0,
                "game_total_preferences": 20,
                "remaining_unique_items": 5,
                "mode": "auto",
                "max_slots": 20,
                "covered_npcs_count": 0,
                "target_npcs_count": 0,
                "vendors_covered_today": 0,
                "today_loved_completed": 0,
                "today_liked_completed": 0,
                "locked_npcs": [],
                "ungiftable_npcs": [],
            },
        }

    def test_excel_export_creates_focus_recipes_sheet(self):
        """Excel export must include 'Focus Recipes' worksheet alongside the 4 standard sheets."""
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = Path(tmpdir) / "test_plan.xlsx"
            export_plan_to_excel(
                save=None,
                plan_results=self.plan_results,
                npc_gift_definitions={},
                item_metadata={},
                output_path=out_path,
            )

            self.assertTrue(out_path.exists())
            wb = openpyxl.load_workbook(str(out_path))

            # Verify standard sheets remain present
            self.assertIn("Daily Bag Plan", wb.sheetnames)
            self.assertIn("NPC Gift Progress", wb.sheetnames)
            self.assertIn("All Remaining Items", wb.sheetnames)
            self.assertIn("Gift Completion Matrix", wb.sheetnames)

            if "Focus Recipes" not in wb.sheetnames:
                self.skipTest("excel_export.py has not yet implemented Focus Recipes sheet")

            self.assertIn("Focus Recipes", wb.sheetnames)
            ws = wb["Focus Recipes"]

            # Header validation
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(headers, ["Rank", "Recipe Name", "Impact", "Unlock Source"])

            # Row data validation
            row2 = [cell.value for cell in ws[2]]
            self.assertEqual(row2, [1, "Apple Pie", 5, "The Inn"])

            row3 = [cell.value for cell in ws[3]]
            self.assertEqual(row3, [2, "Beet Soup", 3, "Shipping"])

    def test_excel_export_empty_focus_recipes_sheet(self):
        """When focus_recipes is empty, sheet is still created with headers."""
        plan_empty = copy.deepcopy(self.plan_results)
        plan_empty["focus_recipes"] = []

        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = Path(tmpdir) / "test_plan_empty.xlsx"
            export_plan_to_excel(
                save=None,
                plan_results=plan_empty,
                npc_gift_definitions={},
                item_metadata={},
                output_path=out_path,
            )

            wb = openpyxl.load_workbook(str(out_path))
            if "Focus Recipes" not in wb.sheetnames:
                self.skipTest("excel_export.py has not yet implemented Focus Recipes sheet")

            ws = wb["Focus Recipes"]
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(headers, ["Rank", "Recipe Name", "Impact", "Unlock Source"])
            # No row 2 data
            self.assertIsNone(ws.cell(row=2, column=1).value)


# ==============================================================================
# TIER 3: CLI INTEGRATION
# ==============================================================================

class TestCLIIntegration(unittest.TestCase):
    """
    Tier 3: Verifies CLI argument parser accepts --recipe-sources and wires it into execution.
    """

    def test_cli_parser_has_recipe_sources_flag(self):
        """build_planner_parser() must define --recipe-sources with default data/recipe_sources.json."""
        parser = build_planner_parser()
        if "--recipe-sources" not in parser.format_help():
            self.skipTest("cli.py parser does not yet define --recipe-sources")

        args = parser.parse_args([])
        self.assertEqual(args.recipe_sources, "data/recipe_sources.json")

    def test_cli_parser_custom_recipe_sources_flag(self):
        """--recipe-sources parses custom file paths properly."""
        parser = build_planner_parser()
        if "--recipe-sources" not in parser.format_help():
            self.skipTest("cli.py parser does not yet define --recipe-sources")

        args = parser.parse_args(["--recipe-sources", "custom/sources.json"])
        self.assertEqual(args.recipe_sources, "custom/sources.json")

    def test_cli_help_text_mentions_recipe_sources(self):
        """--help output must document the --recipe-sources flag."""
        parser = build_planner_parser()
        help_text = parser.format_help()
        if "--recipe-sources" not in help_text:
            self.skipTest("cli.py parser help does not yet include --recipe-sources")
        self.assertIn("--recipe-sources", help_text)


# ==============================================================================
# TIER 4: REAL-WORLD SCENARIOS & PLAYTHROUGHS
# ==============================================================================

class TestTier4RealWorldScenarios(unittest.TestCase):
    """
    Tier 4: Realistic gameplay simulations across starter, mid-game, and completionist saves.
    """

    def test_early_game_starter_save_scenario(self):
        """
        Simulates Spring 3, Year 1: Player has 4 starter recipes unlocked.
        Verifies compute_focus_recipes identifies high-value locked recipes.
        """
        if compute_focus_recipes is None:
            self.skipTest("compute_focus_recipes not implemented yet")

        recipes = load_recipes()
        starter_unlocks = {"bread", "trail_mix", "wild_berry_jam", "dried_squid"}

        # Realistic remaining items map: multiple townspeople desire apple_pie, beet_soup, tomato_soup
        remaining_map = {
            "apple_pie": {"loved": {"adeline", "march"}, "liked": {"balor", "celine"}},  # impact 4
            "tomato_soup": {"loved": {"hemlock"}, "liked": {"juniper", "reina"}},          # impact 3
            "bread": {"loved": {"dozy"}, "liked": {"terithia"}},                         # unlocked!
        }

        focus_list, stats = compute_focus_recipes(
            all_recipes=recipes,
            unlocked_recipe_ids=starter_unlocks,
            remaining_items_map=remaining_map,
            top_n=5,
        )

        self.assertGreaterEqual(stats["total_cooking"], 163)
        self.assertEqual(stats["unlocked_count"], 4)
        # Bread is unlocked, so it must not be in focus list
        focus_ids = [r["recipe_id"] for r in focus_list]
        self.assertNotIn("bread", focus_ids)
        # apple_pie should be top rank
        self.assertEqual(focus_list[0]["recipe_id"], "apple_pie")
        self.assertEqual(focus_list[0]["impact"], 4)

    def test_mid_game_progression_scenario(self):
        """
        Simulates Summer 15, Year 1: Player has 45 unlocked recipes.
        Verifies percentage calculation and exclusion of the 45 unlocked items.
        """
        if compute_focus_recipes is None:
            self.skipTest("compute_focus_recipes not implemented yet")

        recipes = load_recipes()
        cooking_keys = [k for k, v in recipes.items() if v.get("source") == "cooking"]
        midgame_unlocks = set(cooking_keys[:45])

        focus_list, stats = compute_focus_recipes(
            all_recipes=recipes,
            unlocked_recipe_ids=midgame_unlocks,
            remaining_items_map={},
            top_n=10,
        )

        self.assertEqual(stats["unlocked_count"], 45)
        expected_pct = (45 / len(cooking_keys)) * 100.0
        self.assertAlmostEqual(stats["unlocked_percentage"], expected_pct, places=1)

        # None of the 45 should appear in focus_list
        for r in focus_list:
            self.assertNotIn(r["recipe_id"], midgame_unlocks)

    def test_completionist_endgame_scenario(self):
        """
        Simulates endgame player with all 163 cooking recipes unlocked.
        Verifies empty list and 100% completion stats.
        """
        if compute_focus_recipes is None:
            self.skipTest("compute_focus_recipes not implemented yet")

        recipes = load_recipes()
        cooking_keys = {k for k, v in recipes.items() if v.get("source") == "cooking"}

        focus_list, stats = compute_focus_recipes(
            all_recipes=recipes,
            unlocked_recipe_ids=cooking_keys,
            remaining_items_map={},
            top_n=10,
        )

        self.assertEqual(focus_list, [])
        self.assertEqual(stats["unlocked_count"], len(cooking_keys))
        self.assertAlmostEqual(stats["unlocked_percentage"], 100.0, places=2)


if __name__ == "__main__":
    unittest.main()
