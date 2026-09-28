#!/usr/bin/env python3
"""
tests/test_alt_sources.py

Comprehensive Test Suite for Alternate Acquisition Sources & Crafting Dependency Tree (Requirement R6).
Covers:
- TestAltSourcesDataLoader:
  - test_load_alt_sources_default: loads >= 150 items and >= 200 sources
  - test_alt_sources_schema_and_types: validates 12 source types, 15 location icon keys, cost/currency types
  - test_load_alt_sources_missing_file: returns {} gracefully
  - test_load_alt_sources_corrupted_json: returns {} gracefully
  - test_load_alt_sources_normalization: keys are stripped lowercase strings
  - test_alias_keys: checks milk/cow_milk, golden_milk/golden_cow_milk, wood/basic_wood, stone/ore_stone
  - test_load_alt_sources_explicit_path: handles explicit str and Path paths
  - test_load_alt_sources_non_dict_json: returns {} gracefully on non-dict root JSON
- TestCraftingTreeEngine:
  - test_build_focus_crafting_trees_basic: blocker raw materials trace to downstream gift products
  - test_build_focus_crafting_trees_cycle_safety: circular recipes (A -> B -> A) terminate safely
  - test_build_focus_crafting_trees_diamond_dag: diamond DAG (Root -> B -> D and Root -> C -> D) preserved
  - test_build_focus_crafting_trees_prunes_unrelated_recipes: recipes without pending gifts pruned
  - test_build_focus_crafting_trees_empty_inputs: handles empty/None arguments gracefully
  - test_node_annotations: checks item_id, item_name, alt_sources, gift_npcs, children, sprite_url
  - test_plan_daily_gift_bag_integration: verifies focus_trees and alt_sources in plan_daily_gift_bag()
  - test_plan_max_relationship_integration: verifies focus_trees and alt_sources in plan_max_relationship()
  - test_build_focus_crafting_trees_multiple_blockers: handles multiple focus suggestions
  - test_build_focus_crafting_trees_deduplicate_children: deduplicates identical child products under parent
- TestPresentationSurfaces:
  - test_terminal_tree_rendering: box-drawing characters, emoji badges, leaf recipient labels
  - test_terminal_tree_rendering_missing_focus_trees: graceful fallback when focus_trees is empty or absent
  - test_json_exporter_focus_trees: plan_to_json() serializes focus_trees with proper node schema
  - test_server_location_sprite_endpoint: GET /assets/sprites/locations/general_store (200 PNG) and nonexistent (404)
  - test_excel_focus_trees_sheet: exports Sheet 6 "Focus Trees" with headers and depth indentation
  - test_excel_focus_trees_sheet_empty_trees: exports Sheet 6 headers gracefully when focus_trees is empty
"""

import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import openpyxl
from starlette.testclient import TestClient

from companion.config import CompanionConfig
from companion.json_exporter import plan_to_json
from companion.server import app
from fom_planner.crafting import load_recipes
from fom_planner.data_loader import (
    load_alt_sources,
    load_item_metadata,
    load_npc_preferences_from_json,
    load_recipe_sources,
)
from fom_planner.exporters.excel_export import export_plan_to_excel
from fom_planner.exporters.terminal import print_terminal_plan
from fom_planner.models import SaveData
from fom_planner.optimizer import (
    build_focus_crafting_trees,
    plan_daily_gift_bag,
    plan_max_relationship,
)
from tests.test_gift_planner import create_synthetic_save_entries

VALID_SOURCE_TYPES = {
    "shop",
    "inn",
    "market_stall",
    "chicken_statue",
    "mimic",
    "mill",
    "fishing",
    "wishing_well",
    "festival",
    "date",
    "quest",
    "museum",
}

VALID_LOCATION_ICONS = {
    "balor_wagon",
    "chicken_statue",
    "date",
    "festival_stall",
    "fishing",
    "general_store",
    "inn",
    "mill",
    "mimic",
    "museum",
    "quest_board",
    "ranch_shop",
    "saturday_market",
    "tackle_shop",
    "wishing_well",
}


class TestAltSourcesDataLoader(unittest.TestCase):
    """Test suite for Requirement R1 & R2: data layer and load_alt_sources() loader."""

    def test_load_alt_sources_default(self):
        """Verify load_alt_sources() loads default data with 150+ items and 200+ sources."""
        data = load_alt_sources()
        self.assertIsInstance(data, dict)
        self.assertGreaterEqual(
            len(data),
            150,
            f"Expected at least 150 distinct items in alt_sources.json, found {len(data)}",
        )
        total_sources = sum(len(sources) for sources in data.values())
        self.assertGreaterEqual(
            total_sources,
            200,
            f"Expected at least 200 source entries in alt_sources.json, found {total_sources}",
        )

    def test_alt_sources_schema_and_types(self):
        """Verify all sources match the 12 defined types, 15 location icon keys, and cost/currency constraints."""
        data = load_alt_sources()
        observed_types = set()

        for item_id, sources in data.items():
            self.assertEqual(item_id, item_id.strip().lower(), f"Key '{item_id}' not normalized")
            self.assertIsInstance(sources, list, f"Sources for '{item_id}' must be a list")
            self.assertGreater(len(sources), 0, f"Sources for '{item_id}' must not be empty")

            for source in sources:
                self.assertIsInstance(source, dict, f"Source entry for '{item_id}' must be a dict")
                self.assertIn("type", source, f"Source entry for '{item_id}' missing 'type'")
                stype = source["type"]
                self.assertIn(
                    stype,
                    VALID_SOURCE_TYPES,
                    f"Invalid source type '{stype}' for item '{item_id}'",
                )
                observed_types.add(stype)

                # Icon key validation
                if source.get("icon") is not None:
                    self.assertIn(
                        source["icon"],
                        VALID_LOCATION_ICONS,
                        f"Unknown icon '{source['icon']}' for item '{item_id}'",
                    )

                # Cost and currency validation
                if source.get("cost") is not None:
                    self.assertIsInstance(
                        source["cost"],
                        int,
                        f"Cost must be integer for item '{item_id}'",
                    )
                    self.assertGreaterEqual(
                        source["cost"],
                        0,
                        f"Cost must be non-negative for item '{item_id}'",
                    )

                if source.get("currency") is not None:
                    self.assertIn(
                        source["currency"],
                        {"tesserae", "shiny_beads", "beads", "t"},
                        f"Unexpected currency '{source['currency']}' for item '{item_id}'",
                    )

                # Optional string fields
                for str_field in ("vendor", "location", "note"):
                    if source.get(str_field) is not None:
                        self.assertIsInstance(
                            source[str_field],
                            str,
                            f"Field '{str_field}' must be a string for item '{item_id}'",
                        )

        # All 12 source types must be represented in the dataset
        self.assertEqual(
            observed_types,
            VALID_SOURCE_TYPES,
            f"Missing source types: {VALID_SOURCE_TYPES - observed_types}",
        )

    def test_load_alt_sources_missing_file(self):
        """Verify load_alt_sources() returns {} gracefully when file does not exist."""
        nonexistent = Path("/nonexistent/dir/alt_sources_missing_xyz.json")
        result = load_alt_sources(nonexistent)
        self.assertEqual(result, {})

    def test_load_alt_sources_corrupted_json(self):
        """Verify load_alt_sources() returns {} gracefully when JSON is malformed."""
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as tf:
            tf.write("{malformed json syntax, not valid")
            temp_name = tf.name
        try:
            result = load_alt_sources(temp_name)
            self.assertEqual(result, {})
        finally:
            if os.path.exists(temp_name):
                os.unlink(temp_name)

    def test_load_alt_sources_normalization(self):
        """Verify keys are normalized to stripped lowercase strings."""
        sample = {
            "  FLOUR  ": [{"type": "shop", "vendor": "General Store"}],
            "Sugar\n": [{"type": "shop", "vendor": "General Store"}],
            "  cOW_mILK  ": [{"type": "shop", "vendor": "Balor's Wagon"}],
        }
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as tf:
            json.dump(sample, tf)
            temp_name = tf.name
        try:
            result = load_alt_sources(temp_name)
            self.assertIn("flour", result)
            self.assertIn("sugar", result)
            self.assertIn("cow_milk", result)
            self.assertNotIn("  FLOUR  ", result)
            self.assertNotIn("Sugar\n", result)
            self.assertNotIn("  cOW_mILK  ", result)
            self.assertEqual(len(result["flour"]), 1)
        finally:
            if os.path.exists(temp_name):
                os.unlink(temp_name)

    def test_alias_keys(self):
        """Verify canonical and alias keys are present and consistent in alt_sources.json."""
        data = load_alt_sources()
        alias_pairs = [
            ("milk", "cow_milk"),
            ("golden_milk", "golden_cow_milk"),
            ("wood", "basic_wood"),
            ("stone", "ore_stone"),
        ]
        for colloquial, canonical in alias_pairs:
            self.assertIn(
                colloquial,
                data,
                f"Colloquial alias '{colloquial}' missing from alt_sources.json",
            )
            self.assertIn(
                canonical,
                data,
                f"Canonical key '{canonical}' missing from alt_sources.json",
            )
            self.assertGreater(len(data[colloquial]), 0)
            self.assertGreater(len(data[canonical]), 0)
            colloquial_types = {s["type"] for s in data[colloquial]}
            canonical_types = {s["type"] for s in data[canonical]}
            self.assertEqual(
                colloquial_types,
                canonical_types,
                f"Source types mismatch between '{colloquial}' and '{canonical}'",
            )

    def test_load_alt_sources_explicit_path(self):
        """Verify load_alt_sources() accepts explicit str or Path path."""
        canonical_path = REPO_ROOT / "data" / "alt_sources.json"
        res_from_path = load_alt_sources(canonical_path)
        res_from_str = load_alt_sources(str(canonical_path))
        self.assertGreaterEqual(len(res_from_path), 150)
        self.assertEqual(len(res_from_path), len(res_from_str))

    def test_load_alt_sources_non_dict_json(self):
        """Verify load_alt_sources() returns {} gracefully when JSON root is not a dictionary."""
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as tf:
            json.dump(["array", "of", "strings"], tf)
            temp_name = tf.name
        try:
            result = load_alt_sources(temp_name)
            self.assertEqual(result, {})
        finally:
            if os.path.exists(temp_name):
                os.unlink(temp_name)


class TestCraftingTreeEngine(unittest.TestCase):
    """Test suite for Requirement R2: build_focus_crafting_trees() DAG traversal & optimizer integration."""

    def test_build_focus_crafting_trees_basic(self):
        """Verify downstream tree construction from a blocker raw material to intermediate products and final gifts."""
        focus_suggestions = [{"item_id": "cow_milk", "item_name": "Milk"}]
        recipes = {
            "cheese_recipe": {
                "item_id": "cheese",
                "display_name": "Cheese",
                "ingredients": [{"item_id": "cow_milk", "count": 1}],
            },
            "cheesecake_recipe": {
                "item_id": "cheesecake",
                "display_name": "Cheesecake",
                "ingredients": [{"item_id": "cheese", "count": 1}],
            },
        }
        remaining_items_map = {
            "cheesecake": {"loved": {"darcy", "eiland"}},
        }
        alt_sources = {
            "cow_milk": [{"type": "shop", "vendor": "Balor's Wagon"}],
            "cheese": [{"type": "chicken_statue", "cost": 10, "currency": "shiny_beads"}],
            "cheesecake": [{"type": "chicken_statue", "cost": 100, "currency": "shiny_beads"}],
        }

        trees = build_focus_crafting_trees(
            focus_suggestions=focus_suggestions,
            recipes=recipes,
            remaining_items_map=remaining_items_map,
            alt_sources=alt_sources,
        )

        self.assertEqual(len(trees), 1)
        root = trees[0]
        self.assertEqual(root["item_id"], "cow_milk")
        self.assertEqual(root["item_name"], "Milk")
        self.assertEqual(len(root["alt_sources"]), 1)
        self.assertEqual(len(root["children"]), 1)

        child_cheese = root["children"][0]
        self.assertEqual(child_cheese["item_id"], "cheese")
        self.assertEqual(child_cheese["item_name"], "Cheese")
        self.assertEqual(len(child_cheese["children"]), 1)

        leaf_cheesecake = child_cheese["children"][0]
        self.assertEqual(leaf_cheesecake["item_id"], "cheesecake")
        self.assertEqual(leaf_cheesecake["item_name"], "Cheesecake")
        self.assertEqual(sorted(leaf_cheesecake["gift_npcs"]), ["Darcy", "Eiland"])
        self.assertEqual(len(leaf_cheesecake["children"]), 0)

    def test_build_focus_crafting_trees_cycle_safety(self):
        """Verify circular recipes (A -> B -> A) terminate safely without infinite recursion or RecursionError."""
        focus_suggestions = [{"item_id": "item_a", "item_name": "Item A"}]
        recipes = {
            "rec_b": {
                "item_id": "item_b",
                "ingredients": [{"item_id": "item_a", "count": 1}],
            },
            "rec_a": {
                "item_id": "item_a",
                "ingredients": [{"item_id": "item_b", "count": 1}],
            },
        }
        remaining_items_map = {
            "item_b": {"loved": {"adeline"}},
        }

        trees = build_focus_crafting_trees(
            focus_suggestions=focus_suggestions,
            recipes=recipes,
            remaining_items_map=remaining_items_map,
        )

        self.assertEqual(len(trees), 1)
        root = trees[0]
        self.assertEqual(root["item_id"], "item_a")
        self.assertEqual(len(root["children"]), 1)
        child_b = root["children"][0]
        self.assertEqual(child_b["item_id"], "item_b")
        # Child b should terminate and not recursively include item_a
        self.assertEqual(len(child_b["children"]), 0)

    def test_build_focus_crafting_trees_diamond_dag(self):
        """Verify diamond dependencies (Root -> B -> D and Root -> C -> D) are properly preserved."""
        focus_suggestions = [{"item_id": "root_item", "item_name": "Root"}]
        recipes = {
            "rec_b": {
                "item_id": "branch_b",
                "ingredients": [{"item_id": "root_item", "count": 1}],
            },
            "rec_c": {
                "item_id": "branch_c",
                "ingredients": [{"item_id": "root_item", "count": 1}],
            },
            "rec_d_from_b": {
                "item_id": "target_d",
                "ingredients": [{"item_id": "branch_b", "count": 1}],
            },
            "rec_d_from_c": {
                "item_id": "target_d",
                "ingredients": [{"item_id": "branch_c", "count": 1}],
            },
        }
        remaining_items_map = {
            "target_d": {"loved": {"celine"}},
        }

        trees = build_focus_crafting_trees(
            focus_suggestions=focus_suggestions,
            recipes=recipes,
            remaining_items_map=remaining_items_map,
        )

        self.assertEqual(len(trees), 1)
        root = trees[0]
        child_ids = [c["item_id"] for c in root["children"]]
        self.assertIn("branch_b", child_ids)
        self.assertIn("branch_c", child_ids)

        b_node = next(c for c in root["children"] if c["item_id"] == "branch_b")
        c_node = next(c for c in root["children"] if c["item_id"] == "branch_c")

        self.assertTrue(any(sub["item_id"] == "target_d" for sub in b_node["children"]))
        self.assertTrue(any(sub["item_id"] == "target_d" for sub in c_node["children"]))

    def test_build_focus_crafting_trees_prunes_unrelated_recipes(self):
        """Verify recipes that do not lead to any pending gift in remaining_items_map are pruned."""
        focus_suggestions = [{"item_id": "root", "item_name": "Root"}]
        recipes = {
            "useful_rec": {
                "item_id": "useful_product",
                "ingredients": [{"item_id": "root", "count": 1}],
            },
            "unrelated_rec": {
                "item_id": "unrelated_product",
                "ingredients": [{"item_id": "root", "count": 1}],
            },
            "unrelated_chain": {
                "item_id": "deep_unrelated",
                "ingredients": [{"item_id": "unrelated_product", "count": 1}],
            },
        }
        # Only useful_product is wanted
        remaining_items_map = {
            "useful_product": {"loved": {"hayden"}},
        }

        trees = build_focus_crafting_trees(
            focus_suggestions=focus_suggestions,
            recipes=recipes,
            remaining_items_map=remaining_items_map,
        )

        self.assertEqual(len(trees), 1)
        root = trees[0]
        child_ids = [c["item_id"] for c in root["children"]]
        self.assertEqual(child_ids, ["useful_product"])
        self.assertNotIn("unrelated_product", child_ids)

    def test_build_focus_crafting_trees_empty_inputs(self):
        """Verify empty suggestions, empty recipes, or empty remaining_items_map handle gracefully."""
        # Empty focus suggestions
        self.assertEqual(build_focus_crafting_trees([], {"a": {}}, {"a": {}}), [])
        # None focus suggestions
        self.assertEqual(build_focus_crafting_trees(None, {"a": {}}, {"a": {}}), [])

        # Empty recipes: produces root node with empty children
        trees_no_recipes = build_focus_crafting_trees(
            [{"item_id": "milk"}], {}, {"milk": {"loved": {"darcy"}}}
        )
        self.assertEqual(len(trees_no_recipes), 1)
        self.assertEqual(trees_no_recipes[0]["children"], [])

        # Empty remaining items map: all downstream branches pruned, children empty
        recipes = {"cheese": {"item_id": "cheese", "ingredients": [{"item_id": "milk", "count": 1}]}}
        trees_no_remaining = build_focus_crafting_trees([{"item_id": "milk"}], recipes, {})
        self.assertEqual(len(trees_no_remaining), 1)
        self.assertEqual(trees_no_remaining[0]["children"], [])

        # None inputs
        self.assertEqual(build_focus_crafting_trees([], None, None, None, None), [])

    def test_node_annotations(self):
        """Verify node schema: item_id, item_name, alt_sources, gift_npcs, children, and sprite_url."""
        focus_suggestions = [{"item_id": "flour"}]
        recipes = {
            "bread": {
                "item_id": "bread",
                "display_name": "Fresh Bread",
                "ingredients": [{"item_id": "flour", "count": 1}],
            }
        }
        remaining_items_map = {"bread": {"loved": {"celine"}}}
        alt_sources = {
            "flour": [{"type": "shop", "vendor": "General Store", "cost": 200, "currency": "tesserae", "icon": "general_store"}],
            "bread": [{"type": "quest", "location": "Request Board", "note": "Reward for Celine's Request", "icon": "quest_board"}],
        }
        item_metadata = {
            "flour": {"display_name": "Wheat Flour"},
            "bread": {"display_name": "Fresh Bread"},
        }

        trees = build_focus_crafting_trees(
            focus_suggestions=focus_suggestions,
            recipes=recipes,
            remaining_items_map=remaining_items_map,
            alt_sources=alt_sources,
            item_metadata=item_metadata,
        )

        self.assertEqual(len(trees), 1)
        root = trees[0]
        for field in ("item_id", "item_name", "alt_sources", "gift_npcs", "children", "sprite_url"):
            self.assertIn(field, root)

        self.assertEqual(root["item_id"], "flour")
        self.assertEqual(root["item_name"], "Wheat Flour")
        self.assertEqual(root["sprite_url"], "/assets/sprites/items/flour")
        self.assertEqual(len(root["alt_sources"]), 1)
        self.assertEqual(root["alt_sources"][0]["vendor"], "General Store")

        self.assertEqual(len(root["children"]), 1)
        child = root["children"][0]
        for field in ("item_id", "item_name", "alt_sources", "gift_npcs", "children", "sprite_url"):
            self.assertIn(field, child)

        self.assertEqual(child["item_id"], "bread")
        self.assertEqual(child["item_name"], "Fresh Bread")
        self.assertEqual(child["sprite_url"], "/assets/sprites/items/bread")
        self.assertEqual(child["gift_npcs"], ["Celine"])
        self.assertEqual(len(child["alt_sources"]), 1)
        self.assertEqual(child["alt_sources"][0]["type"], "quest")

    def test_plan_daily_gift_bag_integration(self):
        """Verify plan_daily_gift_bag() returns 'focus_trees' and 'alt_sources'."""
        npcs, _ = load_npc_preferences_from_json(REPO_ROOT / "data" / "item_data.json")
        meta = load_item_metadata(REPO_ROOT / "assets" / "fiddle", REPO_ROOT / "data" / "item_data.json")
        recipes = load_recipes()
        sources = load_recipe_sources()
        raw = create_synthetic_save_entries(bag_items={"bread": 2, "apple": 5})
        save = SaveData(Path("test.sav"), raw)

        plan = plan_daily_gift_bag(
            save=save,
            npc_gift_definitions=npcs,
            item_metadata=meta,
            recipe_sources=sources,
        )

        self.assertIn("focus_trees", plan)
        self.assertIn("alt_sources", plan)
        self.assertIsInstance(plan["focus_trees"], list)
        self.assertIsInstance(plan["alt_sources"], dict)
        self.assertGreaterEqual(len(plan["alt_sources"]), 150)

        if plan["focus_trees"]:
            root = plan["focus_trees"][0]
            for field in ("item_id", "item_name", "alt_sources", "gift_npcs", "children", "sprite_url"):
                self.assertIn(field, root)

    def test_plan_max_relationship_integration(self):
        """Verify plan_max_relationship() returns 'focus_trees' and 'alt_sources'."""
        npcs, _ = load_npc_preferences_from_json(REPO_ROOT / "data" / "item_data.json")
        meta = load_item_metadata(REPO_ROOT / "assets" / "fiddle", REPO_ROOT / "data" / "item_data.json")
        recipes = load_recipes()
        sources = load_recipe_sources()
        raw = create_synthetic_save_entries(bag_items={"bread": 2, "apple": 5})
        save = SaveData(Path("test.sav"), raw)

        plan = plan_max_relationship(
            save=save,
            npc_gift_definitions=npcs,
            item_metadata=meta,
            recipe_sources=sources,
        )

        self.assertIn("focus_trees", plan)
        self.assertIn("alt_sources", plan)
        self.assertIsInstance(plan["focus_trees"], list)
        self.assertIsInstance(plan["alt_sources"], dict)
        self.assertGreaterEqual(len(plan["alt_sources"]), 150)

    def test_build_focus_crafting_trees_multiple_blockers(self):
        """Verify multiple focus suggestions produce multiple distinct trees."""
        focus_suggestions = [
            {"item_id": "cow_milk", "item_name": "Milk"},
            {"item_id": "sugar", "item_name": "Sugar"},
        ]
        recipes = {
            "cheese": {
                "item_id": "cheese",
                "ingredients": [{"item_id": "cow_milk", "count": 1}],
            },
            "cookies": {
                "item_id": "cookies",
                "ingredients": [{"item_id": "sugar", "count": 1}],
            },
        }
        remaining_items_map = {
            "cheese": {"loved": {"darcy"}},
            "cookies": {"loved": {"ryis"}},
        }
        trees = build_focus_crafting_trees(focus_suggestions, recipes, remaining_items_map)
        self.assertEqual(len(trees), 2)
        self.assertEqual(trees[0]["item_id"], "cow_milk")
        self.assertEqual(trees[1]["item_id"], "sugar")
        self.assertEqual(trees[0]["children"][0]["item_id"], "cheese")
        self.assertEqual(trees[1]["children"][0]["item_id"], "cookies")

    def test_build_focus_crafting_trees_deduplicate_children(self):
        """Verify multiple recipes producing the same product from an ingredient do not duplicate children."""
        focus_suggestions = [{"item_id": "raw_ore", "item_name": "Raw Ore"}]
        recipes = {
            "smelt_1": {
                "item_id": "ingot",
                "ingredients": [{"item_id": "raw_ore", "count": 1}],
            },
            "smelt_bulk": {
                "item_id": "ingot",
                "ingredients": [{"item_id": "raw_ore", "count": 5}],
            },
        }
        remaining_items_map = {"ingot": {"loved": {"march"}}}
        trees = build_focus_crafting_trees(focus_suggestions, recipes, remaining_items_map)
        self.assertEqual(len(trees), 1)
        # Should have exactly 1 child for ingot, not 2
        self.assertEqual(len(trees[0]["children"]), 1)
        self.assertEqual(trees[0]["children"][0]["item_id"], "ingot")


class TestPresentationSurfaces(unittest.TestCase):
    """Test suite for Requirements R3, R4, and R5: Terminal, Web/JSON/Sprite, and Excel outputs."""

    def test_terminal_tree_rendering(self):
        """Verify print_terminal_plan formats box-drawing characters, emoji badges, and leaf recipients."""
        raw = create_synthetic_save_entries(bag_items={"flour": 1})
        save = SaveData(Path("test.sav"), raw)
        plan = {
            "bag_plan": [],
            "npc_progress": {},
            "overall_stats": {
                "today_loved_completed": 0,
                "today_liked_completed": 0,
                "today_points_earned": 0,
                "vendor_deliveries_count": 0,
                "covered_npcs_count": 0,
                "target_npcs_count": 0,
            },
            "focus_suggestions": [
                {
                    "item_id": "milk",
                    "item_name": "Milk",
                    "count": 5,
                    "count_needed": 5,
                    "deficit": 5,
                    "location": "Ranch (Cows)",
                    "blocked_npcs": ["Darcy"],
                    "impact_score": 10,
                }
            ],
            "focus_trees": [
                {
                    "item_id": "milk",
                    "item_name": "Milk",
                    "alt_sources": [
                        {"type": "shop", "vendor": "General Store", "cost": 200, "currency": "tesserae"},
                        {"type": "chicken_statue", "cost": 10, "currency": "shiny_beads"},
                    ],
                    "gift_npcs": [],
                    "children": [
                        {
                            "item_id": "cheese",
                            "item_name": "Cheese",
                            "alt_sources": [
                                {"type": "chicken_statue", "cost": 10, "currency": "shiny_beads"}
                            ],
                            "gift_npcs": [],
                            "children": [
                                {
                                    "item_id": "cheesecake",
                                    "item_name": "Cheesecake",
                                    "alt_sources": [
                                        {"type": "chicken_statue", "cost": 100, "currency": "shiny_beads"}
                                    ],
                                    "gift_npcs": ["Darcy"],
                                    "children": [],
                                }
                            ],
                        }
                    ],
                }
            ],
        }

        buf = io.StringIO()
        with patch("sys.stdout", buf):
            print_terminal_plan(save, plan)
        out = buf.getvalue()

        self.assertIn("🌳 Crafting Tree:", out)
        self.assertTrue("├── " in out or "└── " in out)
        self.assertIn("General Store", out)
        self.assertIn("Cheesecake", out)
        self.assertIn("→ Darcy", out)
        self.assertTrue(any(e in out for e in ["🛒", "🐔", "🍽️"]))

    def test_terminal_tree_rendering_missing_focus_trees(self):
        """Verify print_terminal_plan renders gracefully without crashing when focus_trees is empty or omitted."""
        raw = create_synthetic_save_entries(bag_items={"flour": 1})
        save = SaveData(Path("test.sav"), raw)
        plan = {
            "bag_plan": [],
            "npc_progress": {},
            "overall_stats": {
                "today_loved_completed": 0,
                "today_liked_completed": 0,
                "today_points_earned": 0,
                "vendor_deliveries_count": 0,
                "covered_npcs_count": 0,
                "target_npcs_count": 0,
            },
            "focus_suggestions": [
                {
                    "item_id": "milk",
                    "item_name": "Milk",
                    "count": 5,
                    "count_needed": 5,
                    "deficit": 5,
                    "location": "Ranch (Cows)",
                    "blocked_npcs": ["Darcy"],
                    "impact_score": 10,
                }
            ],
            "focus_trees": [],
        }

        buf = io.StringIO()
        with patch("sys.stdout", buf):
            print_terminal_plan(save, plan)
        out = buf.getvalue()
        self.assertIn("Milk", out)
        self.assertNotIn("🌳 Crafting Tree:", out)

    def test_json_exporter_focus_trees(self):
        """Verify plan_to_json() includes serialized focus_trees in JSON payload."""
        raw = create_synthetic_save_entries(bag_items={"flour": 1})
        save = SaveData(Path("test.sav"), raw)
        config = CompanionConfig()
        plan_results = {
            "bag_plan": [],
            "npc_progress": {},
            "overall_stats": {},
            "focus_suggestions": [],
            "focus_trees": [
                {
                    "item_id": "cow_milk",
                    "item_name": "Milk",
                    "alt_sources": [{"type": "shop", "vendor": "General Store", "cost": 200, "currency": "tesserae"}],
                    "gift_npcs": [],
                    "sprite_url": "/assets/sprites/items/cow_milk",
                    "children": [
                        {
                            "item_id": "cheese",
                            "item_name": "Cheese",
                            "alt_sources": [],
                            "gift_npcs": ["Darcy"],
                            "sprite_url": "/assets/sprites/items/cheese",
                            "children": [],
                        }
                    ],
                }
            ],
        }

        json_dict = plan_to_json(save, plan_results, {}, Path("test.sav"), config)
        self.assertIn("focus_trees", json_dict)
        self.assertEqual(len(json_dict["focus_trees"]), 1)

        # Confirm JSON serialization
        serialized = json.dumps(json_dict)
        self.assertIn('"focus_trees"', serialized)

        deserialized = json.loads(serialized)
        tree = deserialized["focus_trees"][0]
        self.assertEqual(tree["item_id"], "cow_milk")
        self.assertEqual(tree["item_name"], "Milk")
        self.assertEqual(tree["sprite_url"], "/assets/sprites/items/cow_milk")
        self.assertEqual(len(tree["children"]), 1)
        self.assertEqual(tree["children"][0]["gift_npcs"], ["Darcy"])

    def test_server_location_sprite_endpoint(self):
        """Verify /assets/sprites/locations/{name} returns 200 PNG for existing icon and 404 for nonexistent icon."""
        with TestClient(app) as client:
            res_200 = client.get("/assets/sprites/locations/general_store")
            self.assertEqual(res_200.status_code, 200)
            self.assertEqual(res_200.headers.get("content-type"), "image/png")

            res_404 = client.get("/assets/sprites/locations/nonexistent_location_icon_xyz")
            self.assertEqual(res_404.status_code, 404)

    def test_excel_focus_trees_sheet(self):
        """Verify export_plan_to_excel exports Sheet 6 'Focus Trees' with headers and depth indentation."""
        raw = create_synthetic_save_entries(bag_items={"flour": 1})
        save = SaveData(Path("test.sav"), raw)
        plan_results = {
            "bag_plan": [],
            "npc_progress": {},
            "remaining_items_map": {},
            "overall_stats": {},
            "focus_suggestions": [],
            "focus_trees": [
                {
                    "item_id": "milk",
                    "item_name": "Milk",
                    "alt_sources": [{"type": "shop", "vendor": "General Store", "cost": 200, "currency": "tesserae"}],
                    "gift_npcs": [],
                    "children": [
                        {
                            "item_id": "cheese",
                            "item_name": "Cheese",
                            "alt_sources": [{"type": "chicken_statue", "cost": 10, "currency": "shiny_beads"}],
                            "gift_npcs": [],
                            "children": [
                                {
                                    "item_id": "cheesecake",
                                    "item_name": "Cheesecake",
                                    "alt_sources": [{"type": "chicken_statue", "cost": 100, "currency": "shiny_beads"}],
                                    "gift_npcs": ["Darcy", "Eiland"],
                                    "children": [],
                                }
                            ],
                        }
                    ],
                }
            ],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "test_report.xlsx"
            export_plan_to_excel(save, plan_results, {}, {}, out_file)
            self.assertTrue(out_file.exists())

            wb = openpyxl.load_workbook(str(out_file))
            self.assertIn("Focus Trees", wb.sheetnames)
            self.assertEqual(wb.sheetnames[5], "Focus Trees")

            ws = wb["Focus Trees"]
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(
                headers,
                ["Blocker Item", "Depth", "Product", "Alt Sources", "Gift For (NPCs)"],
            )

            data_rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(data_rows), 3)

            # Depth 0: Root Blocker
            self.assertEqual(data_rows[0][0], "Milk")
            self.assertEqual(data_rows[0][1], 0)
            self.assertEqual(data_rows[0][2], "Milk")
            self.assertIn("General Store", data_rows[0][3])
            self.assertEqual(data_rows[0][4], "—")

            # Depth 1: Intermediate Cheese
            self.assertEqual(data_rows[1][0], "Milk")
            self.assertEqual(data_rows[1][1], 1)
            self.assertEqual(data_rows[1][2], "→ Cheese")
            self.assertIn("Chicken Statue", data_rows[1][3])
            self.assertEqual(data_rows[1][4], "—")

            # Depth 2: Final Gift Cheesecake
            self.assertEqual(data_rows[2][0], "Milk")
            self.assertEqual(data_rows[2][1], 2)
            self.assertEqual(data_rows[2][2], "→ → Cheesecake")
            self.assertIn("Chicken Statue", data_rows[2][3])
            self.assertEqual(data_rows[2][4], "Darcy, Eiland")

    def test_excel_focus_trees_sheet_empty_trees(self):
        """Verify export_plan_to_excel creates Sheet 6 'Focus Trees' with headers even when focus_trees is empty."""
        raw = create_synthetic_save_entries(bag_items={"flour": 1})
        save = SaveData(Path("test.sav"), raw)
        plan_results = {
            "bag_plan": [],
            "npc_progress": {},
            "remaining_items_map": {},
            "overall_stats": {},
            "focus_suggestions": [],
            "focus_trees": [],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            out_file = Path(tmpdir) / "empty_trees_report.xlsx"
            export_plan_to_excel(save, plan_results, {}, {}, out_file)
            wb = openpyxl.load_workbook(str(out_file))
            self.assertIn("Focus Trees", wb.sheetnames)
            ws = wb["Focus Trees"]
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(
                headers,
                ["Blocker Item", "Depth", "Product", "Alt Sources", "Gift For (NPCs)"],
            )
            data_rows = list(ws.iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(data_rows), 0)


if __name__ == "__main__":
    unittest.main()
