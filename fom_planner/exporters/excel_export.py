"""
excel_export.py

Multi-sheet Excel workbook export routines and styling using openpyxl.
"""

import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from fom_planner.constants import SATURDAY_MARKET_VENDORS
from fom_planner.models import InGameDate, SaveData

def get_excel_styles():
    """Returns standard openpyxl style objects for workbook formatting."""
    if not OPENPYXL_AVAILABLE:
        return {}
    return {
        "header_fill": PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid"),
        "header_font": Font(name="Segoe UI", size=11, bold=True, color="FFFFFF"),
        "sub_header_fill": PatternFill(start_color="34495E", end_color="34495E", fill_type="solid"),
        "love_fill": PatternFill(start_color="FFD1DC", end_color="FFD1DC", fill_type="solid"),  # Soft pink
        "love_font": Font(name="Segoe UI", size=10, bold=True, color="900C3F"),
        "like_fill": PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid"),  # Soft green
        "like_font": Font(name="Segoe UI", size=10, bold=True, color="1E8449"),
        "done_fill": PatternFill(start_color="EAECEE", end_color="EAECEE", fill_type="solid"),  # Soft gray
        "done_font": Font(name="Segoe UI", size=10, italic=True, color="7F8C8D"),
        "accent_fill": PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid"),  # Soft blue
        "accent_font": Font(name="Segoe UI", size=10, bold=True, color="1B4F72"),
        "data_font": Font(name="Segoe UI", size=10),
        "bold_font": Font(name="Segoe UI", size=10, bold=True),
        "center_align": Alignment(horizontal="center", vertical="center"),
        "left_align": Alignment(horizontal="left", vertical="center"),
        "right_align": Alignment(horizontal="right", vertical="center"),
        "thin_border": Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0")
        )
    }


def style_sheet_table(ws, headers, data_rows, numeric_cols=None, center_cols=None, styles=None):
    """Applies clean, standard styling to an openpyxl worksheet table."""
    if styles is None:
        styles = get_excel_styles()

    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = styles["header_fill"]
        cell.font = styles["header_font"]
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, r_data in enumerate(data_rows, start=2):
        ws.append(r_data)
        ws.row_dimensions[row_idx].height = 20
        for col_idx in range(1, len(r_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = styles["data_font"]
            cell.border = styles["thin_border"]
            if center_cols and col_idx in center_cols:
                cell.alignment = styles["center_align"]
            elif numeric_cols and col_idx in numeric_cols:
                cell.alignment = styles["right_align"]
            else:
                cell.alignment = styles["left_align"]

    # Auto column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions




def export_plan_to_excel(
    save: Optional[SaveData],
    plan_results: dict,
    npc_gift_definitions: Dict[str, dict],
    item_metadata: Dict[str, dict],
    output_path: Path
):
    """
    Exports a comprehensive multi-sheet Excel report with Saturday Market classification and inventory status.
    """
    if not OPENPYXL_AVAILABLE:
        print("[Excel] openpyxl is not installed. Skipping Excel export.", file=sys.stderr)
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet
    styles = get_excel_styles()

    # Add extra style for vendor header
    vendor_header_fill = PatternFill(start_color="8E44AD", end_color="8E44AD", fill_type="solid")  # Purple
    vendor_header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    bag_plan = plan_results["bag_plan"]
    npc_progress = plan_results["npc_progress"]
    stats = plan_results["overall_stats"]

    # ----------------------------------------------------
    # Sheet 1: Daily Bag Plan
    # ----------------------------------------------------
    ws1 = wb.create_sheet(title="Daily Bag Plan")
    headers1 = [
        "Slot #", "Pack Qty", "Availability", "Item Name", "Item ID", "Crafting Chain",
        "Market Vendors Covered", "Today Loved Targets", "Today Liked Targets",
        "All Potential Loved", "All Potential Liked", "Bin Price", "Store Price", "Tags", "Description"
    ]
    rows1 = []
    for b in bag_plan:
        rows1.append([
            b["slot"],
            b["quantity_to_pack"],
            b.get("status", ""),
            b["item_name"],
            b["item_id"],
            b.get("crafting_chain", "") or "—",
            ", ".join(b["market_vendors_covered"]) or "—",
            ", ".join(b["loved_recipients_today"]),
            ", ".join(b["liked_recipients_today"]),
            ", ".join(b["all_potential_loved"]),
            ", ".join(b["all_potential_liked"]),
            b["bin_price"],
            b["store_price"],
            b["tags"],
            b["description"]
        ])

    style_sheet_table(ws1, headers1, rows1, numeric_cols=[1, 2, 12, 13], center_cols=[1, 2, 3], styles=styles)

    # ----------------------------------------------------
    # Sheet 2: NPC Gift Progress
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="NPC Gift Progress")
    headers2 = [
        "NPC Name", "Category", "Present Today?", "Heart Points", "Gifted Today?", "Assigned Item (Today)", "Pref Met (Today)",
        "Gifts Given", "Remaining Loved", "Remaining Liked", "Total Remaining", "% Loved Done", "% Total Done"
    ]
    rows2 = []
    for nid in sorted(npc_gift_definitions.keys(), key=lambda x: npc_gift_definitions[x].get("name", x)):
        p = npc_progress[nid]
        category = "Saturday Market Vendor" if p["is_vendor"] else "Townsfolk"
        present_str = "Yes (In Town)" if p["is_present_today"] else "No (Visiting Aldaria)"
        gifted_today_str = "No (Eligible)" if p["can_gift_today"] else "YES (Already Gifted)"
        rows2.append([
            p["name"],
            category,
            present_str,
            f"{p['hp']:.0f}",
            gifted_today_str,
            p["assigned_item_name"] or "—",
            p["assigned_pref_type"] or "—",
            p["gifts_given_count"],
            p["remaining_loved_count"],
            p["remaining_liked_count"],
            p["total_remaining"],
            f"{p['pct_loved_done']:.1f}%",
            f"{p['pct_total_done']:.1f}%"
        ])

    style_sheet_table(ws2, headers2, rows2, numeric_cols=[4, 8, 9, 10, 11, 12, 13], center_cols=[2, 3, 4, 5, 7, 8, 9, 10, 11, 12, 13], styles=styles)

    # ----------------------------------------------------
    # Sheet 3: All Remaining Items
    # ----------------------------------------------------
    ws3 = wb.create_sheet(title="All Remaining Items")
    headers3 = ["Rank", "Item Name", "Item ID", "Remaining Total NPCs", "Market Vendors Need", "Remaining Loved", "Remaining Liked", "Remaining Loved By", "Remaining Liked By", "Bin Price", "Store Price", "Tags"]

    rem_items = plan_results["remaining_items_map"]
    rem_rows = []
    for item_id, targets in rem_items.items():
        meta = item_metadata.get(item_id, {})
        disp_name = meta.get("display_name") or item_id.replace("_", " ").title()
        l_names = sorted([npc_progress[x]["name"] for x in targets["loved"]])
        k_names = sorted([npc_progress[x]["name"] for x in targets["liked"]])
        tot = len(l_names) + len(k_names)
        vendor_names = sorted([npc_progress[x]["name"] for x in (targets["loved"] | targets["liked"]) if npc_progress[x]["is_vendor"]])
        tag_list = meta.get("tags", [])
        tags_str = ", ".join(tag_list) if isinstance(tag_list, list) else str(tag_list)

        rem_rows.append({
            "item_id": item_id,
            "item_name": disp_name,
            "total": tot,
            "vendor_count": len(vendor_names),
            "vendor_names": ", ".join(vendor_names) or "—",
            "loved_count": len(l_names),
            "liked_count": len(k_names),
            "loved_by": ", ".join(l_names),
            "liked_by": ", ".join(k_names),
            "bin_price": meta.get("bin_price", ""),
            "store_price": meta.get("store_price", ""),
            "tags": tags_str,
        })

    rem_rows.sort(key=lambda x: (-x["total"], -x["vendor_count"], -x["loved_count"], -x["liked_count"], x["item_name"].lower()))
    rows3 = []
    for rank, r in enumerate(rem_rows, start=1):
        rows3.append([
            rank, r["item_name"], r["item_id"], r["total"], r["vendor_names"], r["loved_count"], r["liked_count"],
            r["loved_by"], r["liked_by"], r["bin_price"], r["store_price"], r["tags"]
        ])

    style_sheet_table(ws3, headers3, rows3, numeric_cols=[1, 4, 6, 7, 10, 11], center_cols=[1, 4, 6, 7], styles=styles)

    # ----------------------------------------------------
    # Sheet 4: Gift Completion Matrix (Items x 34 NPCs)
    # ----------------------------------------------------
    ws4 = wb.create_sheet(title="Gift Completion Matrix")
    all_npc_keys_sorted = sorted(npc_gift_definitions.keys(), key=lambda x: npc_gift_definitions[x].get("name", x))
    
    matrix_headers = ["Item Name", "Item ID", "Status Today"] + [
        f"{npc_gift_definitions[nid].get('name', nid)} (Market)" if nid.lower() in SATURDAY_MARKET_VENDORS else npc_gift_definitions[nid].get('name', nid)
        for nid in all_npc_keys_sorted
    ]
    ws4.append(matrix_headers)
    ws4.row_dimensions[1].height = 28

    for col_idx in range(1, len(matrix_headers) + 1):
        cell = ws4.cell(row=1, column=col_idx)
        if col_idx > 3 and all_npc_keys_sorted[col_idx - 4].lower() in SATURDAY_MARKET_VENDORS:
            cell.fill = vendor_header_fill
            cell.font = vendor_header_font
        else:
            cell.fill = styles["header_fill"]
            cell.font = styles["header_font"]
        cell.alignment = styles["center_align"]

    all_game_items = set()
    for g in npc_gift_definitions.values():
        all_game_items.update(g.get("loved", []))
        all_game_items.update(g.get("liked", []))

    today_assignments = {}
    for b in bag_plan:
        iid = b["item_id"]
        for r in b["all_recipients_today"]:
            today_assignments[(iid, r["npc_id"])] = "PACK (LOVE)" if r["pref"] == "LOVE" else "PACK (LIKE)"

    sorted_all_items = sorted(list(all_game_items), key=lambda x: item_metadata.get(x, {}).get("display_name", x).lower())

    for row_idx, iid in enumerate(sorted_all_items, start=2):
        meta = item_metadata.get(iid, {})
        disp_name = meta.get("display_name") or iid.replace("_", " ").title()

        in_bag = any(b["item_id"] == iid for b in bag_plan)
        status_today = "IN BAG TODAY" if in_bag else ""

        row_cells = [disp_name, iid, status_today]

        for nid in all_npc_keys_sorted:
            g_def = npc_gift_definitions[nid]
            given_set = save.get_npc_gifts_given(nid) if (save and hasattr(save, "get_npc_gifts_given")) else set()

            if (iid, nid) in today_assignments:
                row_cells.append(today_assignments[(iid, nid)])
            elif iid in given_set:
                row_cells.append("GIFTED")
            elif iid in g_def.get("loved", []):
                row_cells.append("LOVE")
            elif iid in g_def.get("liked", []):
                row_cells.append("LIKE")
            else:
                row_cells.append("")

        ws4.append(row_cells)
        ws4.row_dimensions[row_idx].height = 19

        for col_idx in range(1, len(row_cells) + 1):
            cell = ws4.cell(row=row_idx, column=col_idx)
            cell.border = styles["thin_border"]
            val = str(cell.value or "")

            if col_idx == 1:
                cell.font = styles["bold_font"]
                cell.alignment = styles["left_align"]
            elif col_idx == 2:
                cell.font = styles["data_font"]
                cell.alignment = styles["left_align"]
            elif col_idx == 3:
                cell.font = styles["bold_font"] if in_bag else styles["data_font"]
                cell.alignment = styles["center_align"]
                if in_bag:
                    cell.fill = styles["accent_fill"]
                    cell.font = styles["accent_font"]
            else:
                cell.alignment = styles["center_align"]
                if "PACK" in val:
                    cell.fill = styles["accent_fill"]
                    cell.font = styles["accent_font"]
                elif val == "GIFTED":
                    cell.fill = styles["done_fill"]
                    cell.font = styles["done_font"]
                elif val == "LOVE":
                    cell.fill = styles["love_fill"]
                    cell.font = styles["love_font"]
                elif val == "LIKE":
                    cell.fill = styles["like_fill"]
                    cell.font = styles["like_font"]
                else:
                    cell.font = styles["data_font"]

    ws4.column_dimensions["A"].width = 28
    ws4.column_dimensions["B"].width = 22
    ws4.column_dimensions["C"].width = 16
    for col_idx in range(4, len(matrix_headers) + 1):
        ws4.column_dimensions[get_column_letter(col_idx)].width = 14

    ws4.freeze_panes = "D2"
    ws4.auto_filter.ref = ws4.dimensions

    wb.save(output_path)
    print(f"[Excel] Successfully exported Gift Planner Report to: {output_path}")




def export_to_excel(all_rows: list, npcs: dict, output_path: Path):
    """
    Exports a styled multi-sheet Excel workbook with openpyxl.
    """
    if not OPENPYXL_AVAILABLE:
        print("[Excel] openpyxl is not installed. Skipping Excel export.", file=sys.stderr)
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet
    styles = get_excel_styles()
    header_fill = styles["header_fill"]
    header_font = styles["header_font"]
    center_align = styles["center_align"]
    left_align = styles["left_align"]
    thin_border = styles["thin_border"]
    bold_font = styles["bold_font"]
    data_font = styles["data_font"]
    love_fill = styles["love_fill"]
    love_font = styles["love_font"]
    like_fill = styles["like_fill"]
    like_font = styles["like_font"]

    # ----------------------------------------------------
    # Sheet 1: All Items Ranked
    # ----------------------------------------------------
    ws1 = wb.create_sheet(title="All Items Ranked")
    headers1 = [
        "Rank", "Item Name", "Item ID", "Total NPCs", "Loved Count", "Liked Count",
        "Loved By (NPCs)", "Liked By (NPCs)", "Bin Price", "Store Price", "Tags",
        "Quest Usages", "Cooking Usages", "Crafting Usages"
    ]
    rows1 = [
        [
            r["rank"], r["item_name"], r["item_id"], r["total_count"],
            r["loved_count"], r["liked_count"], r["loved_by"], r["liked_by"],
            r["bin_price"], r["store_price"], r["tags"], r["quest_usage"],
            r["cooking_usage"], r["crafting_usage"]
        ]
        for r in all_rows
    ]
    style_sheet_table(ws1, headers1, rows1, numeric_cols=[1, 4, 5, 6, 9, 10, 12, 13, 14], center_cols=[1, 4, 5, 6])

    # ----------------------------------------------------
    # Sheet 2: Loved Gifts Ranked
    # ----------------------------------------------------
    loved_rows = [r for r in all_rows if r["loved_count"] > 0]
    loved_rows_sorted = sorted(loved_rows, key=lambda x: (-x["loved_count"], -x["total_count"], x["item_name"].lower()))
    ws2 = wb.create_sheet(title="Loved Gifts Ranked")
    headers2 = ["Loved Rank", "Item Name", "Loved Count", "Total NPCs", "Loved By (NPCs)", "Bin Price", "Store Price", "Tags"]
    rows2 = [
        [idx, r["item_name"], r["loved_count"], r["total_count"], r["loved_by"], r["bin_price"], r["store_price"], r["tags"]]
        for idx, r in enumerate(loved_rows_sorted, start=1)
    ]
    style_sheet_table(ws2, headers2, rows2, numeric_cols=[1, 3, 4, 6, 7], center_cols=[1, 3, 4])

    # ----------------------------------------------------
    # Sheet 3: Liked Gifts Ranked
    # ----------------------------------------------------
    liked_rows = [r for r in all_rows if r["liked_count"] > 0]
    liked_rows_sorted = sorted(liked_rows, key=lambda x: (-x["liked_count"], -x["total_count"], x["item_name"].lower()))
    ws3 = wb.create_sheet(title="Liked Gifts Ranked")
    headers3 = ["Liked Rank", "Item Name", "Liked Count", "Total NPCs", "Liked By (NPCs)", "Bin Price", "Store Price", "Tags"]
    rows3 = [
        [idx, r["item_name"], r["liked_count"], r["total_count"], r["liked_by"], r["bin_price"], r["store_price"], r["tags"]]
        for idx, r in enumerate(liked_rows_sorted, start=1)
    ]
    style_sheet_table(ws3, headers3, rows3, numeric_cols=[1, 3, 4, 6, 7], center_cols=[1, 3, 4])

    # ----------------------------------------------------
    # Sheet 4: NPC Gift Matrix (Items x NPCs)
    # ----------------------------------------------------
    ws4 = wb.create_sheet(title="NPC Gift Matrix")
    sorted_npc_ids = sorted(npcs.keys(), key=lambda nid: npcs[nid]["name"])
    matrix_headers = ["Rank", "Item Name", "Total", "Loved", "Liked"] + [npcs[nid]["name"] for nid in sorted_npc_ids]
    
    ws4.append(matrix_headers)
    ws4.row_dimensions[1].height = 28

    for col_idx in range(1, len(matrix_headers) + 1):
        cell = ws4.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for row_idx, r in enumerate(all_rows, start=2):
        row_cells = [r["rank"], r["item_name"], r["total_count"], r["loved_count"], r["liked_count"]]
        loved_ids_set = set(r["loved_by_ids"].split(", ")) if r["loved_by_ids"] else set()
        liked_ids_set = set(r["liked_by_ids"].split(", ")) if r["liked_by_ids"] else set()

        for nid in sorted_npc_ids:
            if nid in loved_ids_set:
                row_cells.append("LOVE")
            elif nid in liked_ids_set:
                row_cells.append("LIKE")
            else:
                row_cells.append("")

        ws4.append(row_cells)
        ws4.row_dimensions[row_idx].height = 19

        # Style row cells
        for col_idx in range(1, len(row_cells) + 1):
            cell = ws4.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            val = cell.value

            if col_idx == 2:
                cell.font = bold_font
                cell.alignment = left_align
            elif col_idx in (1, 3, 4, 5):
                cell.font = data_font
                cell.alignment = center_align
            else:
                cell.alignment = center_align
                if val == "LOVE":
                    cell.fill = love_fill
                    cell.font = love_font
                elif val == "LIKE":
                    cell.fill = like_fill
                    cell.font = like_font
                else:
                    cell.font = data_font

    # Set matrix column widths
    ws4.column_dimensions["A"].width = 7
    ws4.column_dimensions["B"].width = 28
    ws4.column_dimensions["C"].width = 8
    ws4.column_dimensions["D"].width = 8
    ws4.column_dimensions["E"].width = 8
    for col_idx in range(6, len(matrix_headers) + 1):
        ws4.column_dimensions[get_column_letter(col_idx)].width = 11

    ws4.freeze_panes = "C2"
    ws4.auto_filter.ref = ws4.dimensions

    # ----------------------------------------------------
    # Sheet 5: NPC Summary
    # ----------------------------------------------------
    ws5 = wb.create_sheet(title="NPC Summary")
    headers5 = ["NPC Name", "NPC ID", "Total Loved Gifts", "Total Liked Gifts", "Loved Gift Items", "Liked Gift Items"]
    npc_summary_rows = []
    
    # Map item_id to name
    item_name_lookup = {r["item_id"]: r["item_name"] for r in all_rows}

    for nid in sorted_npc_ids:
        n_info = npcs[nid]
        loved_names = sorted([item_name_lookup.get(iid, iid.replace("_", " ").title()) for iid in n_info["loved"]])
        liked_names = sorted([item_name_lookup.get(iid, iid.replace("_", " ").title()) for iid in n_info["liked"]])
        npc_summary_rows.append([
            n_info["name"],
            nid,
            len(loved_names),
            len(liked_names),
            ", ".join(loved_names),
            ", ".join(liked_names)
        ])

    style_sheet_table(ws5, headers5, npc_summary_rows, numeric_cols=[3, 4], center_cols=[2, 3, 4])

    wb.save(output_path)
    print(f"[Excel] Successfully exported multi-sheet workbook to: {output_path}")


