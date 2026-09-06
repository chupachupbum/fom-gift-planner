#!/usr/bin/env python3
"""
gift_planner.py

Daily Bag Loadout Optimizer for Fields of Mistria.
Analyzes player save files (.sav) to detect remaining ungifted items per NPC,
and uses a greedy weighted set cover algorithm with inventory awareness and recursive
crafting calculation to determine the optimal set of items to carry in the player's bag,
maximizing NPC coverage while minimizing inventory slots used.

Special Saturday Market Awareness:
- The 8 visiting vendors (Darcy, Louis, Merri, Stillwell, Taliferro, Vera, Wheedle, Zorel)
  only visit Mistria on Saturdays (days 6, 13, 20, 27).
- On weekdays, the planner automatically focuses on the 26 townsfolk in town.
- On Saturdays (or when using '--mode saturday'), the planner includes all 34 NPCs and
  boosts visiting vendor priority so you never miss their rare weekly visit.

Inventory & Crafting Awareness:
- Scans bag inventory and all farm/world storage chests via SaveData.
- Evaluates candidate item availability using recursive DAG crafting resolution.
- Tiered availability scoring: HAVE (2) > CRAFT (1) > UNAVAILABLE (0).
- Dynamically deducts consumed raw materials upon each gift slot selection.
- Displays availability badges (📦 HAVE, ✅ CRAFT, ❌ NEED) and intermediate recipe chains.

Usage:
    python scripts/gift_planner.py
    python scripts/gift_planner.py --mode saturday
    python scripts/gift_planner.py --mode market-only
    python scripts/gift_planner.py --mode all
    python scripts/gift_planner.py --slots 15
    python scripts/gift_planner.py --format terminal
"""

import argparse
import csv
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

# Ensure stdout supports UTF-8 on Windows
if sys.platform == "win32" and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# Import shared modules
try:
    from save_parser import (
        SaveData,
        InGameDate,
        parse_save_file,
        find_latest_save,
        find_save_files,
        SATURDAY_MARKET_VENDORS,
        ANIMAL_FESTIVAL_ATTENDING_VENDORS,
    )
    from export_gift_rankings import (
        load_npc_preferences_from_fiddle,
        load_npc_preferences_from_json,
        load_item_metadata,
        get_excel_styles,
        style_sheet_table,
        OPENPYXL_AVAILABLE,
    )
    from crafting_calculator import (
        AvailabilityTier,
        CraftingPlan,
        CraftingStep,
        load_recipes,
        evaluate_craftability,
        deduct_crafting_materials,
        format_crafting_chain,
        _calculate_max_craftable,
    )
except ImportError:
    # Fallback to package-relative or parent imports
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    try:
        from scripts.save_parser import (
            SaveData,
            InGameDate,
            parse_save_file,
            find_latest_save,
            find_save_files,
            SATURDAY_MARKET_VENDORS,
            ANIMAL_FESTIVAL_ATTENDING_VENDORS,
        )
        from scripts.export_gift_rankings import (
            load_npc_preferences_from_fiddle,
            load_npc_preferences_from_json,
            load_item_metadata,
            get_excel_styles,
            style_sheet_table,
            OPENPYXL_AVAILABLE,
        )
        from scripts.crafting_calculator import (
            AvailabilityTier,
            CraftingPlan,
            CraftingStep,
            load_recipes,
            evaluate_craftability,
            deduct_crafting_materials,
            format_crafting_chain,
            _calculate_max_craftable,
        )
    except ImportError:
        from save_parser import (
            SaveData,
            InGameDate,
            parse_save_file,
            find_latest_save,
            find_save_files,
            SATURDAY_MARKET_VENDORS,
            ANIMAL_FESTIVAL_ATTENDING_VENDORS,
        )
        from export_gift_rankings import (
            load_npc_preferences_from_fiddle,
            load_npc_preferences_from_json,
            load_item_metadata,
            get_excel_styles,
            style_sheet_table,
            OPENPYXL_AVAILABLE,
        )
        from crafting_calculator import (
            AvailabilityTier,
            CraftingPlan,
            CraftingStep,
            load_recipes,
            evaluate_craftability,
            deduct_crafting_materials,
            format_crafting_chain,
            _calculate_max_craftable,
        )

if OPENPYXL_AVAILABLE:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter


ITEM_LOCATIONS: Dict[str, str] = {
    # Golden animal products (Ranch with high happiness)
    "golden_cow_milk": "Ranch (Cows with high happiness)",
    "golden_milk": "Ranch (Cows with high happiness)",
    "golden_egg": "Ranch (Chickens with high happiness)",
    "golden_duck_egg": "Ranch (Ducks with high happiness)",
    "golden_duck_feather": "Ranch (Ducks with high happiness)",
    "golden_feather": "Ranch (Roosters with high happiness)",
    "golden_sheep_wool": "Ranch (Sheep with high happiness)",
    "golden_alpaca_wool": "Ranch (Alpacas with high happiness)",
    "golden_rabbit_wool": "Ranch (Rabbits with high happiness)",
    "golden_horse_hair": "Ranch (Horses with high happiness)",
    "golden_bull_horn": "Ranch (Bulls with high happiness)",
    "golden_bristle": "Ranch (Capybaras with high happiness)",
    "golden_mayonnaise": "The Mill (Golden Egg)",
    "golden_duck_mayonnaise": "The Mill (Golden Duck Egg)",
    "golden_cheese": "The Mill (Golden Milk)",
    "golden_butter": "The Mill (Golden Milk)",

    # Standard animal products & feed (Ranch / Hayden's Shop)
    "cow_milk": "Ranch (Cows)",
    "milk": "Ranch (Cows)",
    "egg": "Ranch (Chickens)",
    "duck_egg": "Ranch (Ducks)",
    "duck_feather": "Ranch (Ducks)",
    "feather": "Ranch (Roosters)",
    "sheep_wool": "Ranch (Sheep)",
    "alpaca_wool": "Ranch (Alpacas)",
    "rabbit_wool": "Ranch (Rabbits)",
    "horse_hair": "Ranch (Horses)",
    "bull_horn": "Ranch (Bulls)",
    "bristle": "Ranch (Capybaras)",
    "mayonnaise": "The Mill (Chicken Egg)",
    "duck_mayonnaise": "The Mill (Duck Egg)",
    "cheese": "The Mill (Milk)",
    "butter": "The Mill (Milk)",
    "hay": "Hayden's Shop / Cutting Grass",
    "grass_seed": "Hayden's Shop / Cutting Grass",
    "animal_currency": "Ranch / Animal Care (Shiny Beads)",

    # Upper Mines (Floors 1-19)
    "ore_copper": "Upper Mines (Floors 1-19)",
    "copper_ore": "Upper Mines (Floors 1-19)",
    "perfect_copper_ore": "Upper Mines (Floors 1-19)",
    "ore_ruby": "Upper Mines (Floors 1-19)",
    "ruby": "Upper Mines (Floors 1-19)",
    "perfect_ruby": "Upper Mines (Floors 1-19)",
    "copper_ingot": "Blacksmith Smelter (Copper Ore)",
    "upper_mines_mushroom": "The Upper Mines (Floors 1-19)",
    "sweetroot": "The Upper Mines (Floors 1-19)",
    "shadow_flower": "The Upper Mines (Floors 1-19)",
    "narrows_moss": "The Upper Mines (Floors 1-19)",
    "red_toadstool": "The Upper Mines (Floors 2-19)",
    "copper_beetle": "Catching (The Upper Mines Floors 1-19)",
    "cave_shrimp": "The Upper Mines (Floors 1-19)",
    "cave_eel": "Fishing (The Upper Mines, All Seasons)",
    "worm": "The Upper Mines (Floors 1-19 - Breaking Rocks)",
    "miners_helmet": "The Upper Mines (Floors 1-19 - Upper Mines Artifact Set)",

    # Tide Caverns (Floors 21-39)
    "ore_iron": "The Tide Caverns (Floors 21-39)",
    "iron_ore": "The Tide Caverns (Floors 21-39)",
    "perfect_iron_ore": "The Tide Caverns (Floors 21-39)",
    "ore_sapphire": "The Tide Caverns (Floors 21-39)",
    "sapphire": "The Tide Caverns (Floors 21-39)",
    "perfect_sapphire": "The Tide Caverns (Floors 21-39)",
    "iron_ingot": "Blacksmith Smelter (Iron Ore)",
    "jade_dulse": "The Tide Caverns (Floors 21-39)",
    "mines_mussels": "The Tide Caverns (Floors 21-39)",
    "underseaweed": "The Tide Caverns (Floors 21-39)",
    "sea_grapes": "The Tide Caverns (Floors 21-39)",
    "tide_lettuce": "The Tide Caverns (Floors 21-39)",
    "cave_kelp": "The Tide Caverns (Floors 21-39)",
    "cave_mushroom": "The Tide Caverns (Floors 21-39)",
    "wild_mushroom": "The Tide Caverns (Floors 21-39)",
    "archerfish": "Fishing (The Tide Caverns, All Seasons)",
    "sapphire_betta": "Fishing (The Tide Caverns, All Seasons)",
    "coral_mantis": "Catching (The Tide Caverns Floors 21-39)",
    "puddle_spider": "Catching (The Tide Caverns Floors 21-39)",
    "sea_scarab": "Catching (The Tide Caverns Floors 21-39 - Breaking Rocks)",
    "stone_shell": "The Tide Caverns (Floors 21-39 - Tide Cavern Artifact Set)",

    # Deep Earth (Floors 41-59)
    "ore_silver": "The Deep Earth (Floors 41-59)",
    "silver_ore": "The Deep Earth (Floors 41-59)",
    "perfect_silver_ore": "The Deep Earth (Floors 41-59)",
    "ore_emerald": "The Deep Earth (Floors 41-59)",
    "emerald": "The Deep Earth (Floors 41-59)",
    "perfect_emerald": "The Deep Earth (Floors 41-59)",
    "silver_ingot": "Blacksmith Smelter (Silver Ore)",
    "earthshroom": "The Deep Earth (Floors 41-59)",
    "shale_grass": "The Deep Earth (Floors 41-59)",
    "crystal_berries": "The Deep Earth (Floors 41-59)",
    "crystal_rose": "The Deep Earth (Floors 41-59)",
    "rockroot": "The Deep Earth (Floors 41-59)",
    "crystal": "The Deep Earth (Floors 41-59)",
    "glowing_mushroom": "The Deep Earth (Monster Drop - Blue Mushroom)",
    "crystalline_cricket": "Catching (The Deep Earth Floors 41-59)",
    "earth_eel": "Fishing (The Deep Earth, All Seasons)",
    "rock_statue": "The Deep Earth (Floors 41-59 - Deep Earth Artifact Set)",

    # Lava Caves (Floors 61-79)
    "ore_gold": "The Lava Caves (Floors 61-79)",
    "gold_ore": "The Lava Caves (Floors 61-79)",
    "perfect_gold_ore": "The Lava Caves (Floors 61-79)",
    "ore_diamond": "The Lava Caves (Floors 61-79)",
    "diamond": "The Lava Caves (Floors 61-79)",
    "perfect_diamond": "The Lava Caves (Floors 61-79)",
    "gold_ingot": "Blacksmith Smelter (Gold Ore)",
    "ash_mushroom": "The Lava Caves (Floors 61-79)",
    "breath_of_fire": "The Lava Caves (Floors 61-79)",
    "flame_pepper": "The Lava Caves (Floors 61-79)",
    "hot_potato": "The Lava Caves (Floors 61-79)",
    "lava_chestnuts": "The Lava Caves (Floors 61-79)",
    "fire_crystal": "The Lava Caves (Floors 61-79)",
    "obsidian": "The Lava Caves (Floors 61-79)",
    "purple_mushroom": "The Lava Caves (Monster Drop - Purple Mushroom)",
    "lava_piranha": "Fishing (The Lava Caves, All Seasons)",
    "armored_bass": "Fishing (The Lava Caves, All Seasons)",
    "sulfur_crab": "The Lava Caves (Floors 61-79)",
    "diamond_beetle": "Catching (The Lava Caves Floors 61-79)",

    # The Ancient Ruins (Floors 81-99)
    "ore_mistril": "The Ancient Ruins (Floors 81-99)",
    "mistril_ore": "The Ancient Ruins (Floors 81-99)",
    "perfect_mistril_ore": "The Ancient Ruins (Floors 81-99)",
    "ore_pink_diamond": "The Ancient Ruins (Floors 81-99)",
    "pink_diamond": "The Ancient Ruins (Floors 81-99)",
    "perfect_pink_diamond": "The Ancient Ruins (Floors 81-99)",
    "mistril_ingot": "Blacksmith Smelter (Mistril Ore)",
    "spell_fruit": "The Ancient Ruins (Floors 81-99)",
    "chirping_fern": "The Ancient Ruins (Floors 81-99)",
    "void_herb": "The Mines (Any Floor Without a Seal)",
    "essence_blossom": "The Ancient Ruins (Floors 81-99)",
    "ethereal_grass": "The Ancient Ruins (Floors 81-99)",
    "written_root": "The Ancient Ruins (Floors 81-99)",
    "ancient_firefly": "Catching (The Ancient Ruins Floors 81-99)",
    "giant_worm": "Catching (The Ancient Ruins Floors 81-99)",
    "hidden_beetle": "Catching (The Ancient Ruins Floors 81-99)",
    "parchment_moth": "Catching (The Ancient Ruins Floors 81-99)",
    "void_snail": "Catching (The Ancient Ruins Floors 81-99)",
    "luminescent_crab": "The Ancient Ruins (Floors 81-99)",
    "winged_shrimp": "The Ancient Ruins (Floors 81-99)",
    "gazer": "Fishing (The Ancient Ruins, All Seasons)",
    "giant_jellyfish": "Fishing (The Ancient Ruins, All Seasons)",
    "coelacanth": "Fishing (The Ancient Ruins, All Seasons)",
    "dragon_claw": "The Ancient Ruins (Floors 81-99)",
    "dragon_scale": "The Ancient Ruins (Floors 81-99)",
    "dragon_forged_bracelet": "The Ancient Ruins (Floors 81-99)",
    "dragon_pact_tablet": "The Ancient Ruins (Floors 81-99)",
    "hardened_essence": "The Ancient Ruins (Floors 81-99)",
    "dragon_forged_core": "The Dragon Forge (Floors 90)",
    "dragon_forged_fang": "The Dragon Forge (Floors 90)",
    "dragon_forged_horn": "The Dragon Forge (Floors 90)",
    "dragon_forged_powder": "The Dragon Forge (Floors 90)",
    "voidite": "The Ancient Ruins (Floors 81-99)",

    # General Mining, Digging & Monster Drops
    "ore_stone": "The Mines / Farm / Town (Breaking Rocks)",
    "stone": "The Mines / Farm / Town (Breaking Rocks)",
    "refined_stone": "Stone Refinery / Blacksmith Smelter",
    "peat": "Digging / The Mines / Stone Refinery",
    "sod": "Digging / The Mines / Stone Refinery",
    "clay": "Digging / The Mines / Diving",
    "glass": "Blacksmith Smelter (Stone) / The Mines",
    "shards": "Digging / Breaking Pots in The Mines",
    "shard_mass": "Digging / Breaking Pots in The Mines",
    "void_stone": "The Mines (Breaking Rocks / Void Sight)",
    "void_powder": "The Mines (Breaking Rocks / Void Sight)",
    "void_pearl": "Fishing (The Mines Water Ponds)",
    "monster_core": "The Mines (Monster drops)",
    "monster_fang": "The Mines (Monster drops)",
    "monster_horn": "The Mines (Monster drops)",
    "monster_shell": "The Mines (Monster drops)",
    "monster_powder": "The Mines (Monster drops)",
    "monster_wing": "The Mines (Monster drops)",
    "monster_whisker": "The Mines (Monster drops)",
    "monster_block": "The Mines (Monster drops)",
    "ice_block": "Foraging (Winter - Frozen Water / Snow)",
    "stone_horse": "Digging (The Mines - Buried Artifact Set)",

    # Farm Crops & Flowers - Spring
    "turnip": "Farm (Spring Crop)",
    "potato": "Farm (Spring Crop)",
    "cabbage": "Farm (Spring Crop)",
    "carrot": "Farm (Spring Crop)",
    "strawberry": "Farm (Spring Crop)",
    "peas": "Farm (Spring Crop)",
    "pea": "Farm (Spring Crop)",
    "chickpea": "Farm (Spring Crop) / Foraging (Spring - Sweetwater Farm, Narrows)",
    "tulip": "Farm (Spring Flower)",
    "daffodil": "Farm / Foraging (Spring Flower)",
    "lilac": "Farm / Foraging (Spring Flower)",
    "snowdrop_anemone": "Farm / Foraging (Spring Flower)",
    "snowdrop": "Farm / Foraging (Spring Flower)",

    # Farm Crops & Flowers - Summer
    "corn": "Farm (Summer Crop)",
    "tomato": "Farm (Summer Crop)",
    "cucumber": "Farm (Summer Crop)",
    "sugar_cane": "Farm (Summer Crop)",
    "tea": "Farm (Summer Crop)",
    "watermelon": "Farm (Summer Crop)",
    "chili_pepper": "Farm (Summer Crop)",
    "sunflower": "Farm (Summer Flower)",
    "marigold": "Farm (Summer Flower) / Foraging (Summer)",
    "cosmos": "Farm (Summer Flower) / Foraging (Summer - Sweetwater Farm)",
    "daisy": "Farm (Summer Flower) / Foraging (Summer)",
    "iris": "Farm (Summer Flower) / Foraging (Summer)",
    "night_queen": "Farm (Summer Flower) / Foraging (Summer)",
    "catmint": "Farm (Summer Crop) / Foraging (Summer - The Narrows)",
    "basil": "Farm (Summer Crop) / Foraging (Summer - Eastern Road)",
    "dill": "Farm (Summer Crop) / Foraging (Summer)",
    "oregano": "Farm (Summer Crop) / Foraging (Summer - Mistria)",
    "sage": "Farm (Summer Crop) / Foraging (Summer - Narrows, Beach)",
    "thyme": "Farm (Summer Crop) / Foraging (Summer - Sweetwater Farm)",

    # Farm Crops & Flowers - Fall
    "pumpkin": "Farm (Fall Crop)",
    "sweet_potato": "Farm (Fall Crop)",
    "onion": "Farm (Fall Crop)",
    "garlic": "Farm (Fall Crop) / Foraging (Fall - Eastern Road)",
    "broccoli": "Farm (Fall Crop)",
    "cranberry": "Farm (Fall Crop)",
    "wheat": "Farm (Fall Crop)",
    "rice_stalk": "Farm (Fall Crop)",
    "rice": "The Mill (Rice Stalk)",
    "flour": "The Mill (Wheat) / General Store",
    "moon_fruit": "Farm (Fall Crop) / Foraging (Fall - Eastern Road, Sweetwater Farm)",
    "chrysanthemum": "Farm (Fall Flower)",
    "celosia": "Farm (Fall Flower) / Foraging (Fall)",
    "heather": "Farm (Fall Flower) / Foraging (Fall)",
    "viola": "Farm (Fall Flower) / Foraging (Fall - Eastern Road)",
    "rosemary": "Farm (Fall Crop) / Foraging (Fall - Sweetwater Farm)",

    # Farm Crops & Flowers - Winter
    "beet": "Farm (Winter Crop)",
    "cauliflower": "Farm (Winter Crop)",
    "daikon_radish": "Farm (Winter Crop)",
    "radish": "Farm (Winter Crop)",
    "snow_peas": "Farm (Winter Crop)",
    "snow_pea": "Farm (Winter Crop)",
    "burdock_root": "Farm (Winter Crop) / Foraging (Winter - Mistria, Eastern Road)",
    "frost_lily": "Farm (Winter Flower) / Foraging (Winter)",
    "jasmine": "Farm (Winter Flower) / Foraging (Winter)",
    "poinsettia": "Farm (Winter Flower) / Foraging (Winter)",
    "snapdragon": "Farm (Winter Flower) / Foraging (Winter - Western Ruins)",

    # Fruit Trees & Woodcutting
    "apple": "Farm / Orchard (Fall Fruit Tree)",
    "cherry": "Farm / Orchard (Spring Fruit Tree)",
    "orange": "Farm / Orchard (Summer Fruit Tree)",
    "peach": "Farm / Orchard (Summer Fruit Tree)",
    "pear": "Farm / Orchard (Fall Fruit Tree)",
    "pomegranate": "Farm / Orchard (Fall Fruit Tree)",
    "lemon": "Farm / Orchard (Summer Fruit Tree)",
    "coconut": "The Beach (Summer Foraging)",
    "acorn": "Foraging (Oak Trees - All Seasons)",
    "pinecone": "Foraging (Pine Trees - All Seasons)",
    "walnut": "Foraging (The Deep Woods)",
    "basic_wood": "Woodcutting (Trees / Farm / Town)",
    "wood": "Woodcutting (Trees / Farm / Town)",
    "hard_wood": "Woodcutting (Large Stumps / Fallen Logs)",
    "hardwood": "Woodcutting (Large Stumps / Fallen Logs)",
    "honey": "Apiary / Hayden's Shop (Upgraded)",
    "honeycomb": "Chopping Trees / Wild Beehives",
    "fiber": "Cutting Weeds / Grass (All Seasons)",

    # Overworld Foraging
    "dandelion": "Foraging (Spring - Mistria, Narrows, Sweetwater Farm)",
    "morel_mushroom": "Foraging (Spring - Rare Overworld)",
    "plum_blossom": "Foraging (Winter - The Western Ruins)",
    "middlemist": "Foraging (Spring - Rare Overworld)",
    "wild_berries": "Foraging (Spring - Bushes in Mistria, Narrows, Farm)",
    "wild_leek": "Foraging (Spring - Mistria, Narrows, Eastern Road)",
    "fennel": "Foraging (Spring - The Narrows, Eastern Road)",
    "fiddlehead": "Foraging (Spring - The Narrows, Sweetwater Farm)",
    "nettle": "Foraging (Spring - Sweetwater Farm, Eastern Road)",
    "blueberry": "Foraging (Spring - The Eastern Road, Deep Woods)",
    "water_chestnut": "Foraging (Water's Edge - Spring, Summer, Fall)",
    "sesame": "Foraging (Summer - Sweetwater Farm)",
    "hydrangea": "Foraging (Summer - The Eastern Road, Deep Woods)",
    "rose": "The Manor's Gardens (Summer)",
    "wild_grapes": "Foraging (Summer - The Narrows)",
    "blackberry": "Foraging (Fall - Sweetwater Farm, Eastern Road, Narrows)",
    "chestnut": "Foraging (Fall - The Beach, Sweetwater Farm, Narrows)",
    "fog_orchid": "Foraging (Fall - The Eastern Road, Western Ruins)",
    "horseradish": "Foraging (Fall - Sweetwater Farm, Eastern Road, Beach)",
    "glowberry": "Foraging (Winter - Mistria, Sweetwater Farm, Narrows)",
    "oyster_mushroom": "Foraging (Winter - The Eastern Road, Mistria, Beach)",
    "pineshroom": "Foraging (Winter - Sweetwater Farm Pine Trees)",
    "rose_hip": "The Manor's Gardens (Winter)",
    "wintergreen_berry": "Foraging (Winter - The Narrows, Western Ruins, Deep Woods)",
    "crocus": "Foraging (Winter - Eastern Road, Sweetwater Farm, Beach)",
    "holly": "Foraging (Winter - The Eastern Road, Sweetwater Farm)",
    "bell_berry": "Foraging (The Deep Woods - All Seasons)",
    "spirit_mushroom": "Foraging (The Deep Woods - All Seasons)",
    "cattail": "Foraging (Rivers / Ponds - Spring, Summer, Fall)",
    "cattail_fluff": "Foraging (Rivers / Ponds - Spring, Summer, Fall)",
    "temple_flower": "Foraging (The Deep Woods - All Seasons)",

    # Beach Foraging & Diving
    "blue_conch_shell": "The Beach (Foraging - All Seasons)",
    "pink_scallop_shell": "The Beach (Foraging - All Seasons)",
    "sand_dollar": "The Beach (Foraging - All Seasons)",
    "spirula_shell": "The Beach (Foraging - All Seasons)",
    "coral": "The Beach (Foraging / Diving - All Seasons)",
    "seaweed": "The Beach (Foraging / Diving - All Seasons)",
    "clam": "The Beach (Diving - All Seasons)",
    "crab": "The Beach (Diving, Summer)",
    "lobster": "The Beach (Diving, Spring)",
    "shrimp": "The Beach (Diving, Spring)",
    "king_crab": "The Beach (Diving, Winter, Snow/Blizzard)",
    "pearl_clam": "The Beach (Diving, All Seasons)",
    "sea_horse": "The Beach (Diving, Summer)",
    "sea_urchin": "The Beach (Diving, All Seasons)",
    "starfish": "The Beach (Diving, All Seasons)",
    "jellyfish": "The Beach (Diving / Ocean, All Seasons)",
    "octopus": "The Beach (Diving, Fall / Winter)",
    "hermit_crab": "The Beach (Catching / Diving - All Seasons)",

    # Fishing & Aquatic Creatures
    "bonito": "Fishing (Ocean, Summer / Winter)",
    "salmon": "Fishing (River, Spring / Fall)",
    "tuna": "Fishing (Ocean, All Seasons)",
    "red_snapper": "Fishing (Ocean, Spring / Fall)",
    "cod": "Fishing (Ocean, Summer / Winter)",
    "mackerel": "Fishing (Ocean, Spring, Rainy/Thunderstorm)",
    "herring": "Fishing (River, Winter)",
    "sardine": "Fishing (Ocean, All Seasons)",
    "bream": "Fishing (River, Summer)",
    "sea_bream": "Fishing (Ocean, Spring / Fall)",
    "perch": "Fishing (River, Fall, Rainy/Storm)",
    "trout": "Fishing (River, All Seasons)",
    "smallmouth_bass": "Fishing (River, All Seasons)",
    "pike": "Fishing (River, All Seasons)",
    "catfish": "Fishing (Pond, All Seasons)",
    "freshwater_eel": "Fishing (River, Winter)",
    "squid": "Fishing (Ocean, All Seasons)",
    "crayfish": "River (Diving, All Seasons)",
    "alligator_gar": "Fishing (Pond, Winter)",
    "anchovy": "Fishing (Ocean, Spring)",
    "angel_fish": "Fishing (Pond, Spring)",
    "barb": "Fishing (Pond, Spring)",
    "bluefish": "Fishing (Pond, Fall, Rainy/Thunderstorm)",
    "bluegill": "Fishing (River, Spring)",
    "bowfish": "Fishing (River, Winter)",
    "brown_bullhead": "Fishing (Pond, Summer)",
    "brown_trout": "Fishing (Pond, Spring)",
    "bullfrog": "Fishing (Pond, Summer / Winter, Rainy/Snow)",
    "burbot": "Fishing (Pond, Winter, Snow/Blizzard)",
    "carp": "Fishing (Pond, Spring)",
    "chum_salmon": "Fishing (River, Fall)",
    "chum": "Fishing (River, Fall)",
    "electric_catfish": "Fishing (Pond, Summer, Thunderstorm)",
    "flathead_catfish": "Fishing (Pond, Winter)",
    "forest_perch": "Fishing (The Deep Woods, Rainy/Snow)",
    "freshwater_oyster": "River (Diving, All Seasons)",
    "frog": "Pond (Diving, Spring / Summer)",
    "gar": "Fishing (Pond, Winter)",
    "giant_trevally": "Fishing (Ocean, Summer)",
    "goby": "Fishing (River, All Seasons)",
    "golden_trout": "Fishing (River, Spring / Fall)",
    "grayling": "Fishing (River, Winter)",
    "green_sunfish": "Fishing (Pond, Summer)",
    "horse_mackerel": "Fishing (Ocean, Winter)",
    "koi": "Fishing (Pond, Spring / Summer)",
    "lake_trout": "Fishing (The Deep Woods, All Seasons)",
    "largemouth_bass": "Fishing (River, Summer)",
    "mahi_mahi": "Fishing (Ocean, Summer)",
    "massive_minnow": "Fishing (River, All Seasons)",
    "minnow": "Fishing (River, Spring)",
    "moonfish": "Fishing (Ocean, Fall)",
    "muskie": "Fishing (River, Fall)",
    "newt": "Pond (Diving / Fishing, All Seasons)",
    "paper_pond_shell": "Pond (Diving, All Seasons)",
    "paper_pond_snail": "Pond (Diving, All Seasons)",
    "pollock": "Fishing (Ocean, Fall / Winter)",
    "pond_snail": "Pond (Diving, All Seasons)",
    "puffer_fish": "Fishing (Ocean, Summer)",
    "rainbow_trout": "Fishing (Pond, Fall)",
    "river_snail": "River (Diving, All Seasons)",
    "roach": "Fishing (River, Spring / Summer)",
    "rock_bass": "Fishing (River, Spring / Fall)",
    "sea_bass": "Fishing (Ocean, Winter)",
    "shadow_bass": "Fishing (River, Winter)",
    "silver_redhorse": "Fishing (River, Spring / Fall)",
    "snail": "Foraging (All Seasons, Rainy)",
    "snakehead": "Fishing (Pond, Summer)",
    "snapping_turtle": "Pond (Diving, Summer / Fall)",
    "striped_bass": "Fishing (Pond, Fall)",
    "sturgeon": "Fishing (River, Winter)",
    "sunny": "Fishing (Pond, Spring / Summer)",
    "swordfish": "Fishing (Ocean, Summer)",
    "tetra": "Fishing (Pond, Spring)",
    "turtle": "Pond (Diving, Spring / Summer)",
    "walleye": "Fishing (River, Fall)",
    "white_perch": "Fishing (Pond, Fall)",
    "yellow_perch": "Fishing (River, Spring)",

    # Bugs & Insects
    "ant": "Catching (Overworld, Spring / Summer / Fall)",
    "bumblebee": "Catching (Overworld, Spring / Summer / Fall)",
    "butterfly": "Catching (Overworld, Spring)",
    "caterpillar": "Catching (Overworld, Spring)",
    "cicada": "Catching (Overworld, Summer)",
    "cricket": "Catching (Overworld, Summer)",
    "crystal_caterpillar": "Catching (Overworld, Winter)",
    "crystal_wing_moth": "Catching (Overworld, Winter)",
    "dragonfly": "Catching (Overworld, Summer / Fall)",
    "fairy_bee": "Catching (Overworld, Fall)",
    "firefly": "Catching (Overworld, Summer Night)",
    "fuzzy_moth": "Catching (Overworld, All Seasons Night)",
    "grasshopper": "Catching (Overworld, Summer / Fall)",
    "honeybee": "Catching (Overworld, Spring / Summer)",
    "hummingbird_hawk_moth": "Catching (Overworld, Spring / Summer / Fall)",
    "inchworm": "Catching (Overworld, Fall)",
    "jewel_beetle": "Catching (Overworld, Summer)",
    "ladybug": "Catching (Overworld, Spring)",
    "lightning_bug": "Catching (Overworld, Summer Night)",
    "lightning_dragonfly": "Catching (Overworld, Thunderstorm)",
    "luna_moth": "Catching (Overworld, Summer Night)",
    "mantis": "Catching (Overworld, Fall)",
    "mistmoth": "Catching (Overworld, Fall Night)",
    "monarch_butterfly": "Catching (Overworld, Fall)",
    "moth": "Catching (Overworld, Spring / Summer Night)",
    "orchid_mantis": "Catching (Overworld, Spring)",
    "pond_skater": "Catching (Rivers / Ponds - All Seasons)",
    "praying_mantis": "Catching (Overworld, Fall)",
    "rhinoceros_beetle": "Catching (Overworld, Spring / Summer / Fall Night)",
    "roly_poly": "Catching (Overworld, Spring - Breaking Rocks)",
    "scarab_beetle": "Catching (Overworld, Summer)",
    "snail_bug": "Catching (Overworld, Rainy)",
    "snowball_beetle": "Catching (Overworld, Winter Blizzard)",
    "stag_beetle": "Catching (Overworld, Summer)",
    "strobe_firefly": "Catching (Overworld, Summer)",
    "swallowtail_butterfly": "Catching (Overworld, Spring / Summer)",

    # Artifacts & Ancient Relics
    "alda_bronze_sword": "Digging (Sweetwater Farm)",
    "alda_feather_pendant": "Digging (Sweetwater Farm)",
    "alda_gem_bracelet": "Digging (Sweetwater Farm)",
    "alda_clay_pot": "Digging (Sweetwater Farm)",
    "alda_mural_tablet": "Digging (Sweetwater Farm)",
    "aldarian_sword": "Digging (The Narrows)",
    "aldarian_gauntlet": "Digging (The Narrows)",
    "aldarian_war_banner": "Digging (The Narrows)",
    "family_crest_pendant": "Digging (The Narrows - Aldarian Artifact Set)",
    "lost_crown_of_aldaria": "Digging (The Narrows - Aldarian Artifact Set)",
    "ancient_crystal_goblet": "Digging (The Western Ruins)",
    "ancient_gold_coin": "Digging (The Western Ruins)",
    "ancient_horn_circlet": "Digging (The Western Ruins)",
    "ancient_royal_scepter": "Digging (The Western Ruins)",
    "ancient_stone_lantern": "Digging (The Western Ruins)",
    "caldosian_drinking_horn": "Digging (The Eastern Road)",
    "caldosian_sword": "Digging (The Eastern Road)",
    "caldosian_breastplate": "Digging (The Eastern Road)",
    "caldosian_emperor_bust": "Digging (The Eastern Road)",
    "statuette_of_caldarus": "Digging (The Eastern Road)",
    "amber_trapped_insect": "Digging (The Beach)",
    "black_tablet": "Digging (All Dig Spots)",
    "completely_wrong_map": "Digging (All Dig Spots)",
    "muttering_cube": "Digging (All Dig Spots)",
    "unknown_dragon_statuette": "Digging (All Dig Spots)",
    "weightless_stone": "Digging (All Dig Spots)",
    "coin_lump": "Fishing (Water Spots)",
    "crystal_apple": "Digging (The Deep Woods)",
    "petrified_wood": "Digging (The Deep Woods)",
    "meteorite": "Digging (The Farm / Rare Dig Spots)",
    "tiny_dinosaur_skeleton": "Digging (All Dig Spots - Prehistoric Artifact Set)",
    "giant_fish_scale": "Fishing (Aquatic Artifact Set)",
    "rusted_treasure_chest": "Fishing (Aquatic Artifact Set)",
    "rubber_fish": "Fishing (Aquatic Artifact Set)",
    "rock_with_a_hole": "Diving (The Beach / Water - Sunken Artifact Set)",

    # Town Stores, Milling & Inn Drinks
    "sugar": "The Mill (Sugar Cane) / General Store",
    "oil": "General Store",
    "curry_powder": "General Store",
    "soy_sauce": "General Store",
    "chocolate": "General Store",
    "rock_salt": "General Store",
    "beer": "The Sleeping Dragon (Inn)",
    "red_wine": "The Sleeping Dragon (Inn)",
    "white_wine": "The Sleeping Dragon (Inn)",
    "hot_toddy": "The Sleeping Dragon (Inn)",
    "coffee": "The Sleeping Dragon (Inn) / Darcy's Stall",
    "green_tea": "Cooking (Tea)",
    "lavender_tea": "Balor's Wagon / Spring Festival Food Stall",
    "mushroom_brew": "The Sleeping Dragon (Inn) / Balor's Wagon / Darcy's Stall",

}


def resolve_root_raw_materials(
    item_id: str,
    recipes: Dict[str, Any],
    call_stack: Optional[Set[str]] = None,
) -> Dict[str, int]:
    """
    Recursively resolves item_id down its recipe tree to identify all root raw materials
    (items that do not have a recipe). Returns a dict of {raw_item_id: count_needed_per_item}.
    """
    if not item_id:
        return {}
    item_id = str(item_id).strip().lower()
    if not item_id:
        return {}

    if call_stack is None:
        call_stack = set()

    if not recipes or not isinstance(recipes, dict) or item_id in call_stack or item_id not in recipes:
        return {item_id: 1}

    recipe = recipes.get(item_id)
    if not isinstance(recipe, dict):
        return {item_id: 1}

    ingredients = recipe.get("ingredients", [])
    if not isinstance(ingredients, list) or not ingredients:
        return {item_id: 1}

    call_stack.add(item_id)
    try:
        raw_mats: Dict[str, int] = {}
        for ing in ingredients:
            if not isinstance(ing, dict):
                continue
            ing_id = str(ing.get("item_id", "")).strip().lower()
            if not ing_id:
                continue
            try:
                ing_count = max(1, int(round(float(ing.get("count", 1)))))
            except (ValueError, TypeError):
                ing_count = 1
            sub_raw = resolve_root_raw_materials(ing_id, recipes, call_stack)
            for raw_id, count in sub_raw.items():
                raw_mats[raw_id] = raw_mats.get(raw_id, 0) + count * ing_count
        return raw_mats
    finally:
        call_stack.remove(item_id)


def compute_focus_suggestions(
    remaining_items_map: Dict[str, Dict[str, Set[str]]],
    inventory: Optional[Dict[str, int]] = None,
    recipes: Optional[Dict[str, Any]] = None,
    item_metadata: Optional[Dict[str, dict]] = None,
    top_n: int = 5,
) -> List[Dict[str, Any]]:
    """
    Computes a ranked list of top insufficient items (raw materials and direct gifts)
    blocking pending NPC gift opportunities.

    Algorithm:
      1. For each item in remaining_items_map, recursively resolves its recipe tree
         to identify all root raw material ingredients needed.
      2. Sums total demand for each raw ingredient across ALL pending gift-NPC pairs.
      3. Accounts for player's owned inventory, including:
         - Pre-crafted/finished gifts in inventory that directly satisfy pending NPC preferences.
         - Processed/intermediate goods (e.g. Flour for Wheat, Sugar for Sugar Cane, Rice for
           Rice Stalk, Cheese/Butter for Milk) converting their surplus towards root raw material inventory.
      4. Subtracts player's effective inventory pool to find deficit = demand - inventory.
      5. Identifies items where deficit > 0.
      6. Ranks by: (1) blocked NPC-gift pairs count descending,
                   (2) deficit count descending (tiebreaker),
                   (3) display name ascending (tiebreaker).
      7. Returns top N items with 1-based rank, deficit, and location hints.
    """
    if top_n <= 0:
        return []

    if recipes is None:
        recipes = load_recipes()
    if item_metadata is None:
        item_metadata = {}

    clean_inv: Dict[str, int] = {}
    if inventory and isinstance(inventory, dict):
        for k, v in inventory.items():
            if not k:
                continue
            k_clean = str(k).strip().lower()
            if not k_clean:
                continue
            try:
                val = int(round(float(v)))
                if val > 0:
                    clean_inv[k_clean] = clean_inv.get(k_clean, 0) + val
            except (ValueError, TypeError):
                pass

    def _normalize_npc_set(val: Any) -> Set[str]:
        if val is None:
            return set()
        if isinstance(val, (str, bytes)):
            s = str(val).strip()
            return {s} if s else set()
        if hasattr(val, "__iter__"):
            res = set()
            for x in val:
                if x is not None:
                    s = str(x).strip()
                    if s:
                        res.add(s)
            return res
        return set()

    # Identify intermediate items in recipes (crafted items used as ingredients in other recipes)
    all_recipe_ingredients: Set[str] = set()
    if recipes and isinstance(recipes, dict):
        for r in recipes.values():
            if isinstance(r, dict):
                for ing in r.get("ingredients", []):
                    if isinstance(ing, dict) and ing.get("item_id"):
                        all_recipe_ingredients.add(str(ing["item_id"]).strip().lower())
    intermediate_items: Set[str] = (
        {k.strip().lower() for k in recipes.keys()} & all_recipe_ingredients
        if recipes and isinstance(recipes, dict)
        else set()
    )

    # Track direct gift demand for intermediate items to reserve them before converting surplus
    direct_gift_demand: Dict[str, int] = {}
    if remaining_items_map and isinstance(remaining_items_map, dict):
        for item_id, targets in remaining_items_map.items():
            if item_id and isinstance(targets, dict):
                item_clean = str(item_id).strip().lower()
                pending = _normalize_npc_set(targets.get("loved")) | _normalize_npc_set(targets.get("liked"))
                direct_gift_demand[item_clean] = len(pending)

    # Calculate effective inventory for root raw materials (including equivalents from surplus intermediate goods)
    effective_inv: Dict[str, int] = dict(clean_inv)
    for item_id, count in clean_inv.items():
        if item_id in intermediate_items:
            needed_direct = direct_gift_demand.get(item_id, 0)
            surplus_count = max(0, count - needed_direct)
            if surplus_count > 0:
                sub_raws = resolve_root_raw_materials(item_id, recipes)
                for raw_id, raw_cnt in sub_raws.items():
                    if raw_id != item_id:
                        effective_inv[raw_id] = effective_inv.get(raw_id, 0) + raw_cnt * surplus_count

    raw_demand: Dict[str, int] = {}
    raw_blocked_pairs: Dict[str, Set[Tuple[str, str]]] = {}

    if remaining_items_map and isinstance(remaining_items_map, dict):
        for item_id, targets in remaining_items_map.items():
            if not item_id or not isinstance(targets, dict):
                continue
            item_clean = str(item_id).strip().lower()
            if not item_clean:
                continue
            loved_set = _normalize_npc_set(targets.get("loved"))
            liked_set = _normalize_npc_set(targets.get("liked"))
            pending_npcs = loved_set | liked_set
            if not pending_npcs:
                continue

            # If player already owns finished copies of this crafted gift, they don't need raw ingredients
            have_finished = clean_inv.get(item_clean, 0) if (recipes and item_clean in recipes) else 0
            total_needed = len(pending_npcs)
            needed_to_craft = max(0, total_needed - have_finished)

            if not recipes or item_clean not in recipes:
                # Raw item (not craftable)
                raw_demand[item_clean] = raw_demand.get(item_clean, 0) + total_needed
                if item_clean not in raw_blocked_pairs:
                    raw_blocked_pairs[item_clean] = set()
                for nid in pending_npcs:
                    raw_blocked_pairs[item_clean].add((str(nid), item_clean))
            else:
                # Crafted item
                if needed_to_craft > 0:
                    raw_mats = resolve_root_raw_materials(item_clean, recipes)
                    for raw_id, count_per_gift in raw_mats.items():
                        if not raw_id:
                            continue
                        raw_demand[raw_id] = raw_demand.get(raw_id, 0) + count_per_gift * needed_to_craft
                        if raw_id not in raw_blocked_pairs:
                            raw_blocked_pairs[raw_id] = set()
                        for nid in list(pending_npcs)[:needed_to_craft]:
                            raw_blocked_pairs[raw_id].add((str(nid), item_clean))

    suggestions: List[Dict[str, Any]] = []
    for raw_id, total_dem in raw_demand.items():
        inv_count = effective_inv.get(raw_id, 0)
        deficit = total_dem - inv_count
        if deficit > 0:
            blocked_count = len(raw_blocked_pairs.get(raw_id, set()))
            meta = item_metadata.get(raw_id, {}) if isinstance(item_metadata, dict) else {}
            if isinstance(meta, dict):
                disp_name = meta.get("display_name") or raw_id.replace("_", " ").title()
            elif isinstance(meta, str) and meta.strip():
                disp_name = meta.strip()
            else:
                disp_name = raw_id.replace("_", " ").title()
            location = ITEM_LOCATIONS.get(raw_id, "")

            suggestions.append({
                "item_id": raw_id,
                "item_name": disp_name,
                "blocked_pairs": blocked_count,
                "demand": total_dem,
                "inventory": inv_count,
                "deficit": deficit,
                "location_hint": location,
            })

    # Sort by: (1) blocked NPC-gift pairs descending, (2) deficit descending, (3) name ascending
    suggestions.sort(key=lambda x: (-x["blocked_pairs"], -x["deficit"], str(x["item_name"]).lower()))

    # Assign 1-indexed ranks and trim to top_n
    result: List[Dict[str, Any]] = []
    for rank, item in enumerate(suggestions[:top_n], start=1):
        item["rank"] = rank
        result.append(item)

    return result


def plan_daily_gift_bag(
    save: Optional[SaveData] = None,
    npc_gift_definitions: Optional[Dict[str, dict]] = None,
    item_metadata: Optional[Dict[str, dict]] = None,
    mode: str = "auto",
    max_slots: int = 20,
    loved_weight: int = 3,
    liked_weight: int = 1,
    vendor_boost: float = 1.5,
    exclude_npcs: Optional[Set[str]] = None,
    force_all_npcs: bool = False,
    inventory: Optional[Dict[str, int]] = None,
    recipes: Optional[Dict[str, Any]] = None,
) -> dict:
    """
    Computes the optimal bag loadout for today's gifting session with inventory awareness.

    Modes:
      - 'auto' / 'today': Use save calendar date (weekday = 26 townsfolk, Saturday = 34 NPCs)
      - 'saturday' / 'market': Simulate Saturday Market day (all 34 NPCs with vendor boost)
      - 'market-only': Plan only for the 8 visiting Saturday Market vendors
      - 'townsfolk' / 'weekday': Plan only for the 26 permanent townsfolk
      - 'all': Plan for all 34 NPCs regardless of day

    Returns a dict with:
        - bag_plan: list of bag slot dicts (including availability_tier, status, status_badge, crafting_chain)
        - npc_progress: dict of npc_id -> stats
        - covered_npcs: set of npc_ids covered today
        - target_npcs: set of npc_ids included in planning
        - remaining_items_map: dict of ungifted item_id -> {"loved": set(), "liked": set()}
        - focus_suggestions: top 5 insufficient blocker items
        - overall_stats: summary counts
    """
    if exclude_npcs is None:
        exclude_npcs = set()
    if npc_gift_definitions is None:
        npc_gift_definitions = {}
    if item_metadata is None:
        item_metadata = {}
    if recipes is None:
        recipes = load_recipes()

    # Determine current inventory pool (merging bag + chests or using explicit inventory)
    if inventory is not None:
        current_inventory = {}
        for k, v in inventory.items():
            if not k:
                continue
            try:
                cnt = int(round(float(v)))
                if cnt > 0:
                    k_clean = str(k).strip().lower()
                    current_inventory[k_clean] = current_inventory.get(k_clean, 0) + cnt
            except (ValueError, TypeError):
                pass
    elif save is not None and hasattr(save, "get_all_available_items"):
        current_inventory = {}
        for k, v in save.get_all_available_items().items():
            if not k:
                continue
            try:
                cnt = int(round(float(v)))
                if cnt > 0:
                    k_clean = str(k).strip().lower()
                    current_inventory[k_clean] = current_inventory.get(k_clean, 0) + cnt
            except (ValueError, TypeError):
                pass
    else:
        current_inventory = {}

    initial_inventory = dict(current_inventory)

    is_sat = False
    is_animal_fest = False
    if save is not None and hasattr(save, "in_game_date") and save.in_game_date is not None:
        is_sat = getattr(save.in_game_date, "is_saturday", False)
        is_animal_fest = getattr(save.in_game_date, "is_animal_festival", False)

    planning_for_saturday = (mode in ("saturday", "market")) or (mode == "auto" and is_sat)

    # 1. Determine which NPCs are eligible under current mode
    all_npcs = sorted(npc_gift_definitions.keys())
    target_npcs = set()
    ungiftable_npcs = []
    not_present_npcs = []

    npc_progress = {}
    total_game_loved = 0
    total_game_liked = 0
    total_given_loved = 0
    total_given_liked = 0

    for nid in all_npcs:
        if nid in exclude_npcs:
            continue

        g_def = npc_gift_definitions[nid]
        npc_name = g_def.get("name", nid.replace("_", " ").title())
        all_loved = set(g_def.get("loved", []))
        all_liked = set(g_def.get("liked", []))
        is_vendor = nid.lower() in SATURDAY_MARKET_VENDORS

        total_game_loved += len(all_loved)
        total_game_liked += len(all_liked)

        if save is not None and hasattr(save, "get_npc_gifts_given"):
            given_set = save.get_npc_gifts_given(nid)
        else:
            given_set = set()

        given_loved = all_loved & given_set
        given_liked = all_liked & given_set

        total_given_loved += len(given_loved)
        total_given_liked += len(given_liked)

        remaining_loved = all_loved - given_set
        remaining_liked = all_liked - given_set

        # Determine presence based on mode
        if mode in ("saturday", "market"):
            is_present = True
        elif mode == "market-only":
            is_present = is_vendor
        elif mode in ("townsfolk", "weekday"):
            is_present = not is_vendor
        elif mode == "all":
            is_present = True
        else:  # "auto" or "today"
            is_present = save.is_npc_present_in_town_today(nid) if (save and hasattr(save, "is_npc_present_in_town_today")) else True

        if force_all_npcs:
            can_gift = True
        elif mode in ("saturday", "market") and not is_sat:
            can_gift = True
        elif save is not None and hasattr(save, "can_gift_npc_today"):
            can_gift = save.can_gift_npc_today(nid)
        else:
            can_gift = True

        if not is_present:
            not_present_npcs.append(nid)
        elif not can_gift:
            ungiftable_npcs.append(nid)
        elif remaining_loved or remaining_liked:
            target_npcs.add(nid)

        hp = save.get_npc_heart_points(nid) if (save and hasattr(save, "get_npc_heart_points")) else 0.0
        npc_progress[nid] = {
            "name": npc_name,
            "hp": hp,
            "is_vendor": is_vendor,
            "is_present_today": is_present,
            "can_gift_today": can_gift,
            "gifts_given_count": len(given_set),
            "given_loved_count": len(given_loved),
            "given_liked_count": len(given_liked),
            "remaining_loved": remaining_loved,
            "remaining_liked": remaining_liked,
            "remaining_loved_count": len(remaining_loved),
            "remaining_liked_count": len(remaining_liked),
            "total_remaining": len(remaining_loved) + len(remaining_liked),
            "pct_loved_done": (len(given_loved) / len(all_loved) * 100.0) if all_loved else 100.0,
            "pct_total_done": ((len(given_loved) + len(given_liked)) / (len(all_loved) + len(all_liked)) * 100.0) if (all_loved or all_liked) else 100.0,
            "assigned_item_id": None,
            "assigned_item_name": None,
            "assigned_pref_type": None,
        }

    # 2. Build remaining item coverage mapping
    # all_remaining_items_map: All non-excluded NPCs with pending preferences (for focus suggestions and game totals)
    all_remaining_items_map = {}
    for nid, p in npc_progress.items():
        for item_id in p["remaining_loved"]:
            if item_id not in all_remaining_items_map:
                all_remaining_items_map[item_id] = {"loved": set(), "liked": set()}
            all_remaining_items_map[item_id]["loved"].add(nid)

        for item_id in p["remaining_liked"]:
            if item_id not in all_remaining_items_map:
                all_remaining_items_map[item_id] = {"loved": set(), "liked": set()}
            all_remaining_items_map[item_id]["liked"].add(nid)

    # today_remaining_items_map: Items relevant to today's target NPCs (for greedy bag selection)
    today_remaining_items_map = {}
    for nid in target_npcs:
        p = npc_progress[nid]
        for item_id in p["remaining_loved"]:
            if item_id not in today_remaining_items_map:
                today_remaining_items_map[item_id] = {"loved": set(), "liked": set()}
            today_remaining_items_map[item_id]["loved"].add(nid)

        for item_id in p["remaining_liked"]:
            if item_id not in today_remaining_items_map:
                today_remaining_items_map[item_id] = {"loved": set(), "liked": set()}
            today_remaining_items_map[item_id]["liked"].add(nid)

    # 3. Greedy Weighted Set Cover Algorithm with Tiered Availability & Dynamic Pool Deduction
    covered_npcs = set()
    bag_plan = []

    while len(bag_plan) < max_slots and len(covered_npcs) < len(target_npcs):
        best_item_id = None
        best_tuple = None
        best_selected_recipients = []
        best_status_tier = AvailabilityTier.UNAVAILABLE
        best_qty = 0

        for item_id, targets in today_remaining_items_map.items():
            new_loved = (targets["loved"] & target_npcs) - covered_npcs
            new_liked = (targets["liked"] & target_npcs) - covered_npcs

            if not new_loved and not new_liked:
                continue

            max_avail = _calculate_max_craftable(item_id, current_inventory, recipes)
            owned_count = current_inventory.get(item_id, 0)

            # Sort candidate recipients by priority:
            # Loved vendors > Loved townsfolk > Liked vendors > Liked townsfolk
            def recip_sort_key(nid):
                is_v = npc_progress[nid]["is_vendor"]
                is_l = nid in new_loved
                is_boosted = (planning_for_saturday and is_v) or (is_animal_fest and nid.lower() in ANIMAL_FESTIVAL_ATTENDING_VENDORS)
                return (0 if is_l and is_boosted else (1 if is_l else (2 if is_boosted else 3)), npc_progress[nid]["name"])

            sorted_recips = sorted(list(new_loved | new_liked), key=recip_sort_key)

            if max_avail > 0:
                actual_cover_count = min(len(sorted_recips), max_avail)
            else:
                actual_cover_count = len(sorted_recips)

            selected_recips = sorted_recips[:actual_cover_count]
            selected_loved = [nid for nid in selected_recips if nid in new_loved]
            selected_liked = [nid for nid in selected_recips if nid in new_liked]

            # Score calculation with Saturday / Festival vendor boost
            score = 0.0
            for nid in selected_loved:
                weight = float(loved_weight)
                is_boosted = (planning_for_saturday and npc_progress[nid]["is_vendor"]) or (is_animal_fest and nid.lower() in ANIMAL_FESTIVAL_ATTENDING_VENDORS)
                if is_boosted:
                    weight *= vendor_boost
                score += weight

            for nid in selected_liked:
                weight = float(liked_weight)
                is_boosted = (planning_for_saturday and npc_progress[nid]["is_vendor"]) or (is_animal_fest and nid.lower() in ANIMAL_FESTIVAL_ATTENDING_VENDORS)
                if is_boosted:
                    weight *= vendor_boost
                score += weight

            if score <= 0:
                continue

            if max_avail <= 0:
                tier = AvailabilityTier.UNAVAILABLE
            elif owned_count >= actual_cover_count:
                tier = AvailabilityTier.HAVE
            else:
                tier = AvailabilityTier.CRAFT

            meta = item_metadata.get(item_id, {})
            try:
                bin_p = int(meta.get("bin_price")) if meta.get("bin_price") not in ("", None) else 999999
            except (ValueError, TypeError):
                bin_p = 999999

            try:
                store_p = int(meta.get("store_price")) if meta.get("store_price") not in ("", None) else 999999
            except (ValueError, TypeError):
                store_p = 999999

            item_name = str(meta.get("display_name", item_id))

            # Count how many visiting / boosted vendors are covered by this item
            vendor_hits = sum(
                1 for nid in selected_recips
                if (planning_for_saturday and npc_progress[nid]["is_vendor"])
                or (is_animal_fest and nid.lower() in ANIMAL_FESTIVAL_ATTENDING_VENDORS)
            )

            # 7-element Candidate ranking tuple:
            # (availability_tier, score, vendor_hits, len(selected_loved), -bin_p, -store_p, -len(item_name))
            candidate_tuple = (int(tier), score, vendor_hits, len(selected_loved), -bin_p, -store_p, -len(item_name))

            if best_item_id is None or candidate_tuple > best_tuple:
                best_tuple = candidate_tuple
                best_item_id = item_id
                best_selected_recipients = selected_recips
                best_status_tier = tier
                best_qty = actual_cover_count

        if best_item_id is None or best_qty <= 0:
            break

        owned = current_inventory.get(best_item_id, 0)
        best_new_loved = [nid for nid in best_selected_recipients if nid in today_remaining_items_map[best_item_id]["loved"]]
        best_new_liked = [nid for nid in best_selected_recipients if nid in today_remaining_items_map[best_item_id]["liked"]]

        meta = item_metadata.get(best_item_id, {})
        disp_name = meta.get("display_name") or best_item_id.replace("_", " ").title()

        # Build list of recipient info
        recipients = []
        for nid in sorted(best_new_loved):
            recipients.append({
                "npc_id": nid,
                "name": npc_progress[nid]["name"],
                "pref": "LOVE",
                "is_vendor": npc_progress[nid]["is_vendor"]
            })
            npc_progress[nid]["assigned_item_id"] = best_item_id
            npc_progress[nid]["assigned_item_name"] = disp_name
            npc_progress[nid]["assigned_pref_type"] = "LOVE"

        for nid in sorted(best_new_liked):
            recipients.append({
                "npc_id": nid,
                "name": npc_progress[nid]["name"],
                "pref": "LIKE",
                "is_vendor": npc_progress[nid]["is_vendor"]
            })
            npc_progress[nid]["assigned_item_id"] = best_item_id
            npc_progress[nid]["assigned_item_name"] = disp_name
            npc_progress[nid]["assigned_pref_type"] = "LIKE"

        quantity_to_pack = best_qty

        # Evaluate craftability and perform dynamic inventory deduction
        plan = evaluate_craftability(best_item_id, current_inventory, recipes, count_needed=quantity_to_pack)
        current_inventory = deduct_crafting_materials(best_item_id, current_inventory, recipes, count=quantity_to_pack)

        # Availability status string and badge
        if best_status_tier == AvailabilityTier.HAVE:
            status_str = "HAVE"
            status_badge = "📦 HAVE"
        elif best_status_tier == AvailabilityTier.CRAFT:
            status_str = "CRAFT"
            if owned > 0:
                status_badge = f"🔨 HAVE ({owned}) + CRAFT ({quantity_to_pack - owned})"
            else:
                status_badge = "✅ CRAFT"
        else:
            status_str = "UNAVAILABLE"
            status_badge = "❌ NEED"

        crafting_chain = plan.chain_summary if plan is not None else ""

        # Also get all potential NPCs who could take this item across all NPCs in the game
        all_potential_loved = sorted([npc_progress[x]["name"] for x in all_remaining_items_map.get(best_item_id, {}).get("loved", set())])
        all_potential_liked = sorted([npc_progress[x]["name"] for x in all_remaining_items_map.get(best_item_id, {}).get("liked", set())])

        tag_list = meta.get("tags", [])
        tags_str = ", ".join(tag_list) if isinstance(tag_list, list) else str(tag_list)

        bag_plan.append({
            "slot": len(bag_plan) + 1,
            "item_id": best_item_id,
            "item_name": disp_name,
            "quantity_to_pack": quantity_to_pack,
            "status": status_str,
            "availability_tier": int(best_status_tier),
            "status_badge": status_badge,
            "crafting_chain": crafting_chain,
            "crafting_plan": plan,
            "max_craftable": plan.max_craftable if plan is not None else 0,
            "raw_materials_needed": plan.raw_materials_needed if plan is not None else {},
            "loved_recipients_today": [r["name"] for r in recipients if r["pref"] == "LOVE"],
            "liked_recipients_today": [r["name"] for r in recipients if r["pref"] == "LIKE"],
            "all_recipients_today": recipients,
            "market_vendors_covered": [r["name"] for r in recipients if r["is_vendor"]],
            "all_potential_loved": all_potential_loved,
            "all_potential_liked": all_potential_liked,
            "bin_price": meta.get("bin_price", ""),
            "store_price": meta.get("store_price", ""),
            "tags": tags_str,
            "description": meta.get("description", "")
        })

        covered_npcs |= set(best_selected_recipients)

    # Overall statistics
    today_loved_count = sum(len(b["loved_recipients_today"]) for b in bag_plan)
    today_liked_count = sum(len(b["liked_recipients_today"]) for b in bag_plan)
    vendors_covered_today = sum(len(b["market_vendors_covered"]) for b in bag_plan)

    overall_stats = {
        "mode": mode,
        "planning_for_saturday": planning_for_saturday,
        "is_animal_festival": is_animal_fest,
        "target_npcs_count": len(target_npcs),
        "covered_npcs_count": len(covered_npcs),
        "slots_used": len(bag_plan),
        "max_slots": max_slots,
        "today_loved_completed": today_loved_count,
        "today_liked_completed": today_liked_count,
        "today_total_completed": today_loved_count + today_liked_count,
        "vendors_covered_today": vendors_covered_today,
        "game_total_loved": total_game_loved,
        "game_total_liked": total_game_liked,
        "game_total_preferences": total_game_loved + total_game_liked,
        "game_given_loved": total_given_loved,
        "game_given_liked": total_given_liked,
        "game_given_total": total_given_loved + total_given_liked,
        "remaining_unique_items": len(all_remaining_items_map),
        "ungiftable_npcs": [npc_gift_definitions[nid].get("name", nid) for nid in ungiftable_npcs if nid in npc_gift_definitions],
        "not_present_npcs": [npc_gift_definitions[nid].get("name", nid) for nid in not_present_npcs if nid in npc_gift_definitions],
    }

    focus_suggestions = compute_focus_suggestions(
        remaining_items_map=all_remaining_items_map,
        inventory=initial_inventory,
        recipes=recipes,
        item_metadata=item_metadata,
        top_n=5,
    )

    return {
        "bag_plan": bag_plan,
        "npc_progress": npc_progress,
        "covered_npcs": covered_npcs,
        "target_npcs": target_npcs,
        "remaining_items_map": all_remaining_items_map,
        "all_remaining_items_map": all_remaining_items_map,
        "today_remaining_items_map": today_remaining_items_map,
        "focus_suggestions": focus_suggestions,
        "overall_stats": overall_stats,
    }


def print_terminal_plan(save: Optional[SaveData], plan_results: dict):
    """Prints a beautiful, formatted CLI summary of the gift plan with inventory badges and Saturday Market context."""
    stats = plan_results["overall_stats"]
    bag_plan = plan_results["bag_plan"]
    npc_progress = plan_results["npc_progress"]

    player_name = save.player_name if (save and hasattr(save, "player_name")) else "Player"
    farm_name = save.farm_name if (save and hasattr(save, "farm_name")) else "Farm"
    date_info = save.in_game_date if (save and hasattr(save, "in_game_date")) else InGameDate(year=1, season="spring", day=1)
    save_name = save.file_path.name if (save and hasattr(save, "file_path")) else "active_session"

    print("\n" + "═" * 86)
    print(f"  🎁 FIELDS OF MISTRIA — DAILY GIFT BAG PLANNER")
    print(f"  Player: {player_name} @ {farm_name}  |  {date_info}  |  Save: {save_name}")
    print("═" * 86)

    # Storage & inventory status banner
    if save and hasattr(save, "get_all_available_items"):
        bag_items = len(save.get_bag_items()) if hasattr(save, "get_bag_items") else 0
        chest_items = len(save.get_chest_items()) if hasattr(save, "get_chest_items") else 0
        total_avail = len(save.get_all_available_items())
        print(f"\n📦 INVENTORY & STORAGE STATUS:")
        print(f"  • Player Bag: {bag_items} unique items ({save.empty_inventory_slots}/{save.total_inventory_slots} slots empty)")
        print(f"  • Chest Storage: {chest_items} unique items across all farm/world chests")
        print(f"  • Total Available for Gifting/Crafting: {total_avail} unique items")

    # Market & Festival context banner
    if getattr(date_info, "is_saturday", False) or stats["mode"] in ("saturday", "market"):
        if getattr(date_info, "is_saturday", False):
            print(f"\n🎉 TODAY IS SATURDAY MARKET DAY! All 34 NPCs (26 townsfolk + 8 visiting vendors) are in town!")
        else:
            days_until = getattr(date_info, "days_until_saturday", 0)
            next_sat = getattr(date_info, "next_saturday_day", 6)
            print(f"\n🎪 PLANNING FOR SATURDAY MARKET (Simulated Day {next_sat}, in {days_until} days):")
            print(f"  • Includes all 34 NPCs with priority boost on the 8 weekly visiting vendors.")
    elif stats["mode"] == "market-only":
        print(f"\n🎪 PLANNING FOR SATURDAY MARKET VENDORS ONLY:")
        print(f"  • Focuses exclusively on the 8 weekly visiting vendors (Darcy, Louis, Merri, Stillwell, Taliferro, Vera, Wheedle, Zorel).")
    elif getattr(date_info, "is_animal_festival", False) or stats.get("is_animal_festival", False):
        days_until = getattr(date_info, "days_until_saturday", 0)
        next_sat = getattr(date_info, "next_saturday_day", 6)
        print(f"\n🎉 TODAY IS ANIMAL FESTIVAL DAY! (Winter 10)")
        print(f"  • 28 NPCs present in town (26 permanent townsfolk + visiting contestants Merri & Louis)!")
        print(f"  • Merri & Louis receive priority boost for this special non-Saturday festival appearance.")
        aldaria_vendors = sorted([npc_progress[x]["name"] for x in npc_progress if npc_progress[x]["is_vendor"] and x.lower() not in ANIMAL_FESTIVAL_ATTENDING_VENDORS])
        print(f"  ℹ️  6 Saturday Market vendors remain in Aldaria and return this Saturday (Day {next_sat}, in {days_until} days):")
        print(f"     {', '.join(aldaria_vendors)}")
    else:
        day_of_week = getattr(date_info, "day_of_week", "Weekday")
        days_until = getattr(date_info, "days_until_saturday", 0)
        next_sat = getattr(date_info, "next_saturday_day", 6)
        print(f"\n📅 TODAY'S VISITORS ({day_of_week}): 26 Townsfolk in town.")
        vendor_names = sorted([npc_progress[x]["name"] for x in npc_progress if npc_progress[x]["is_vendor"]])
        print(f"  ℹ️  8 Saturday Market vendors are visiting Aldaria and return this Saturday (Day {next_sat}, in {days_until} days):")
        print(f"     {', '.join(vendor_names)}")
        print(f"  💡 Tip: Run with '--mode saturday' to plan your upcoming Saturday Market shopping list!")

    print(f"\n📊 OVERALL COMPLETION PROGRESS:")
    p_loved = (stats["game_given_loved"] / stats["game_total_loved"] * 100) if stats["game_total_loved"] else 0
    p_liked = (stats["game_given_liked"] / stats["game_total_liked"] * 100) if stats["game_total_liked"] else 0
    p_total = (stats["game_given_total"] / stats["game_total_preferences"] * 100) if stats["game_total_preferences"] else 0
    print(f"  • Loved Gifts: {stats['game_given_loved']}/{stats['game_total_loved']} ({p_loved:.1f}%)")
    print(f"  • Liked Gifts: {stats['game_given_liked']}/{stats['game_total_liked']} ({p_liked:.1f}%)")
    print(f"  • Total Preferences: {stats['game_given_total']}/{stats['game_total_preferences']} ({p_total:.1f}%)")
    print(f"  • Unique Ungifted Items Remaining in Game: {stats['remaining_unique_items']}")

    if stats["ungiftable_npcs"]:
        print(f"\n⚠️  ALREADY GIFTED TODAY ({len(stats['ungiftable_npcs'])} NPCs): {', '.join(stats['ungiftable_npcs'])}")

    print(f"\n🎒 OPTIMAL BAG LOADOUT ({len(bag_plan)}/{stats['max_slots']} slots → covers {stats['covered_npcs_count']}/{stats['target_npcs_count']} NPCs):")
    print("━" * 86)
    print(f" {'Slot':<5} │ {'Qty':<4} │ {'Status':<9} │ {'Item Name':<24} │ {'Today Target Recipients (NPCs)'}")
    print("━" * 86)

    for b in bag_plan:
        recipients_display = []
        for r in b["all_recipients_today"]:
            badge = "♥" if r["pref"] == "LOVE" else "♡"
            vendor_tag = " [Market]" if r["is_vendor"] else ""
            recipients_display.append(f"{badge}{r['name']}{vendor_tag}")
        recip_str = ", ".join(recipients_display)

        status_badge = b.get("status_badge", "❌ NEED")
        print(f"  {b['slot']:<4} │ {b['quantity_to_pack']:<4} │ {status_badge:<9} │ {b['item_name']:<24} │ {recip_str}")
        if b.get("crafting_chain"):
            print(f"       │      │           └─ {b['crafting_chain']}")

    print("━" * 86)
    vendor_msg = f" (including {stats['vendors_covered_today']} Saturday Market vendors)" if stats['vendors_covered_today'] else ""
    print(f"🎯 IMPACT: +{stats['today_loved_completed']} Loved & +{stats['today_liked_completed']} Liked preferences completed in one trip!{vendor_msg}")
    if stats['covered_npcs_count'] == stats['target_npcs_count'] and stats['target_npcs_count'] > 0:
        extra_slots = stats['max_slots'] - len(bag_plan)
        print(f"✨ 100% NPC coverage achieved! You have {extra_slots} extra bag slots free for tools/foraging.")

    # Focus Suggestions Section
    focus_suggestions = plan_results.get("focus_suggestions")
    if focus_suggestions is None:
        inv = save.get_all_available_items() if (save and hasattr(save, "get_all_available_items")) else {}
        focus_suggestions = compute_focus_suggestions(
            remaining_items_map=plan_results.get("all_remaining_items_map") or plan_results.get("remaining_items_map", {}),
            inventory=inv,
        )

    print(f"\n💡 FOCUS SUGGESTIONS (Top Blocker Items):")
    print("━" * 86)
    if not focus_suggestions:
        print("  🎉 All required materials and gifts are currently in your inventory!")
    else:
        for s in focus_suggestions:
            loc_str = f"  ({s['location_hint']})" if s.get("location_hint") else ""
            print(f"  {s['rank']}. {s['item_name']} — Need {s['deficit']} more{loc_str}")
    print("━" * 86)
    print("═" * 86 + "\n")


def export_plan_to_csv(plan_results: dict, output_dir: Path, csv_name: str = "daily_gift_bag_plan.csv"):
    """Exports the daily bag plan and NPC progress to CSV files with availability columns."""
    output_dir.mkdir(parents=True, exist_ok=True)
    plan_csv_path = output_dir / csv_name

    bag_plan = plan_results["bag_plan"]
    headers = [
        ("slot", "Slot #"),
        ("quantity_to_pack", "Pack Qty"),
        ("status", "Availability"),
        ("item_name", "Item Name"),
        ("item_id", "Item ID"),
        ("crafting_chain", "Crafting Chain"),
        ("market_vendors_covered", "Market Vendors Covered"),
        ("loved_recipients_today", "Loved Targets (Today)"),
        ("liked_recipients_today", "Liked Targets (Today)"),
        ("bin_price", "Bin Price"),
        ("store_price", "Store Price"),
        ("tags", "Tags"),
        ("description", "Description")
    ]

    with open(plan_csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([h[1] for h in headers])
        for b in bag_plan:
            writer.writerow([
                b["slot"],
                b["quantity_to_pack"],
                b.get("status", ""),
                b["item_name"],
                b["item_id"],
                b.get("crafting_chain", ""),
                ", ".join(b["market_vendors_covered"]),
                ", ".join(b["loved_recipients_today"]),
                ", ".join(b["liked_recipients_today"]),
                b["bin_price"],
                b["store_price"],
                b["tags"],
                b["description"],
            ])

    print(f"[CSV] Successfully exported Daily Bag Plan to: {plan_csv_path}")


def export_plan_to_excel(
    save: Optional[SaveData],
    plan_results: dict,
    npc_gift_definitions: Dict[str, dict],
    item_metadata: Dict[str, dict],
    output_path: Path
):
    """
    Exports a comprehensive multi-sheet Excel report with Saturday Market classification and inventory status.
    """
    if not OPENPYXL_AVAILABLE:
        print("[Excel] openpyxl is not installed. Skipping Excel export.", file=sys.stderr)
        return

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet
    styles = get_excel_styles()

    # Add extra style for vendor header
    vendor_header_fill = PatternFill(start_color="8E44AD", end_color="8E44AD", fill_type="solid")  # Purple
    vendor_header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    bag_plan = plan_results["bag_plan"]
    npc_progress = plan_results["npc_progress"]
    stats = plan_results["overall_stats"]

    # ----------------------------------------------------
    # Sheet 1: Daily Bag Plan
    # ----------------------------------------------------
    ws1 = wb.create_sheet(title="Daily Bag Plan")
    headers1 = [
        "Slot #", "Pack Qty", "Availability", "Item Name", "Item ID", "Crafting Chain",
        "Market Vendors Covered", "Today Loved Targets", "Today Liked Targets",
        "All Potential Loved", "All Potential Liked", "Bin Price", "Store Price", "Tags", "Description"
    ]
    rows1 = []
    for b in bag_plan:
        rows1.append([
            b["slot"],
            b["quantity_to_pack"],
            b.get("status", ""),
            b["item_name"],
            b["item_id"],
            b.get("crafting_chain", "") or "—",
            ", ".join(b["market_vendors_covered"]) or "—",
            ", ".join(b["loved_recipients_today"]),
            ", ".join(b["liked_recipients_today"]),
            ", ".join(b["all_potential_loved"]),
            ", ".join(b["all_potential_liked"]),
            b["bin_price"],
            b["store_price"],
            b["tags"],
            b["description"]
        ])

    style_sheet_table(ws1, headers1, rows1, numeric_cols=[1, 2, 12, 13], center_cols=[1, 2, 3], styles=styles)

    # ----------------------------------------------------
    # Sheet 2: NPC Gift Progress
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="NPC Gift Progress")
    headers2 = [
        "NPC Name", "Category", "Present Today?", "Heart Points", "Gifted Today?", "Assigned Item (Today)", "Pref Met (Today)",
        "Gifts Given", "Remaining Loved", "Remaining Liked", "Total Remaining", "% Loved Done", "% Total Done"
    ]
    rows2 = []
    for nid in sorted(npc_gift_definitions.keys(), key=lambda x: npc_gift_definitions[x].get("name", x)):
        p = npc_progress[nid]
        category = "Saturday Market Vendor" if p["is_vendor"] else "Townsfolk"
        present_str = "Yes (In Town)" if p["is_present_today"] else "No (Visiting Aldaria)"
        gifted_today_str = "No (Eligible)" if p["can_gift_today"] else "YES (Already Gifted)"
        rows2.append([
            p["name"],
            category,
            present_str,
            f"{p['hp']:.0f}",
            gifted_today_str,
            p["assigned_item_name"] or "—",
            p["assigned_pref_type"] or "—",
            p["gifts_given_count"],
            p["remaining_loved_count"],
            p["remaining_liked_count"],
            p["total_remaining"],
            f"{p['pct_loved_done']:.1f}%",
            f"{p['pct_total_done']:.1f}%"
        ])

    style_sheet_table(ws2, headers2, rows2, numeric_cols=[4, 8, 9, 10, 11, 12, 13], center_cols=[2, 3, 4, 5, 7, 8, 9, 10, 11, 12, 13], styles=styles)

    # ----------------------------------------------------
    # Sheet 3: All Remaining Items
    # ----------------------------------------------------
    ws3 = wb.create_sheet(title="All Remaining Items")
    headers3 = ["Rank", "Item Name", "Item ID", "Remaining Total NPCs", "Market Vendors Need", "Remaining Loved", "Remaining Liked", "Remaining Loved By", "Remaining Liked By", "Bin Price", "Store Price", "Tags"]

    rem_items = plan_results["remaining_items_map"]
    rem_rows = []
    for item_id, targets in rem_items.items():
        meta = item_metadata.get(item_id, {})
        disp_name = meta.get("display_name") or item_id.replace("_", " ").title()
        l_names = sorted([npc_progress[x]["name"] for x in targets["loved"]])
        k_names = sorted([npc_progress[x]["name"] for x in targets["liked"]])
        tot = len(l_names) + len(k_names)
        vendor_names = sorted([npc_progress[x]["name"] for x in (targets["loved"] | targets["liked"]) if npc_progress[x]["is_vendor"]])
        tag_list = meta.get("tags", [])
        tags_str = ", ".join(tag_list) if isinstance(tag_list, list) else str(tag_list)

        rem_rows.append({
            "item_id": item_id,
            "item_name": disp_name,
            "total": tot,
            "vendor_count": len(vendor_names),
            "vendor_names": ", ".join(vendor_names) or "—",
            "loved_count": len(l_names),
            "liked_count": len(k_names),
            "loved_by": ", ".join(l_names),
            "liked_by": ", ".join(k_names),
            "bin_price": meta.get("bin_price", ""),
            "store_price": meta.get("store_price", ""),
            "tags": tags_str,
        })

    rem_rows.sort(key=lambda x: (-x["total"], -x["vendor_count"], -x["loved_count"], -x["liked_count"], x["item_name"].lower()))
    rows3 = []
    for rank, r in enumerate(rem_rows, start=1):
        rows3.append([
            rank, r["item_name"], r["item_id"], r["total"], r["vendor_names"], r["loved_count"], r["liked_count"],
            r["loved_by"], r["liked_by"], r["bin_price"], r["store_price"], r["tags"]
        ])

    style_sheet_table(ws3, headers3, rows3, numeric_cols=[1, 4, 6, 7, 10, 11], center_cols=[1, 4, 6, 7], styles=styles)

    # ----------------------------------------------------
    # Sheet 4: Gift Completion Matrix (Items x 34 NPCs)
    # ----------------------------------------------------
    ws4 = wb.create_sheet(title="Gift Completion Matrix")
    all_npc_keys_sorted = sorted(npc_gift_definitions.keys(), key=lambda x: npc_gift_definitions[x].get("name", x))
    
    matrix_headers = ["Item Name", "Item ID", "Status Today"] + [
        f"{npc_gift_definitions[nid].get('name', nid)} (Market)" if nid.lower() in SATURDAY_MARKET_VENDORS else npc_gift_definitions[nid].get('name', nid)
        for nid in all_npc_keys_sorted
    ]
    ws4.append(matrix_headers)
    ws4.row_dimensions[1].height = 28

    for col_idx in range(1, len(matrix_headers) + 1):
        cell = ws4.cell(row=1, column=col_idx)
        if col_idx > 3 and all_npc_keys_sorted[col_idx - 4].lower() in SATURDAY_MARKET_VENDORS:
            cell.fill = vendor_header_fill
            cell.font = vendor_header_font
        else:
            cell.fill = styles["header_fill"]
            cell.font = styles["header_font"]
        cell.alignment = styles["center_align"]

    all_game_items = set()
    for g in npc_gift_definitions.values():
        all_game_items.update(g.get("loved", []))
        all_game_items.update(g.get("liked", []))

    today_assignments = {}
    for b in bag_plan:
        iid = b["item_id"]
        for r in b["all_recipients_today"]:
            today_assignments[(iid, r["npc_id"])] = "PACK (LOVE)" if r["pref"] == "LOVE" else "PACK (LIKE)"

    sorted_all_items = sorted(list(all_game_items), key=lambda x: item_metadata.get(x, {}).get("display_name", x).lower())

    for row_idx, iid in enumerate(sorted_all_items, start=2):
        meta = item_metadata.get(iid, {})
        disp_name = meta.get("display_name") or iid.replace("_", " ").title()

        in_bag = any(b["item_id"] == iid for b in bag_plan)
        status_today = "IN BAG TODAY" if in_bag else ""

        row_cells = [disp_name, iid, status_today]

        for nid in all_npc_keys_sorted:
            g_def = npc_gift_definitions[nid]
            given_set = save.get_npc_gifts_given(nid) if (save and hasattr(save, "get_npc_gifts_given")) else set()

            if (iid, nid) in today_assignments:
                row_cells.append(today_assignments[(iid, nid)])
            elif iid in given_set:
                row_cells.append("GIFTED")
            elif iid in g_def.get("loved", []):
                row_cells.append("LOVE")
            elif iid in g_def.get("liked", []):
                row_cells.append("LIKE")
            else:
                row_cells.append("")

        ws4.append(row_cells)
        ws4.row_dimensions[row_idx].height = 19

        for col_idx in range(1, len(row_cells) + 1):
            cell = ws4.cell(row=row_idx, column=col_idx)
            cell.border = styles["thin_border"]
            val = str(cell.value or "")

            if col_idx == 1:
                cell.font = styles["bold_font"]
                cell.alignment = styles["left_align"]
            elif col_idx == 2:
                cell.font = styles["data_font"]
                cell.alignment = styles["left_align"]
            elif col_idx == 3:
                cell.font = styles["bold_font"] if in_bag else styles["data_font"]
                cell.alignment = styles["center_align"]
                if in_bag:
                    cell.fill = styles["accent_fill"]
                    cell.font = styles["accent_font"]
            else:
                cell.alignment = styles["center_align"]
                if "PACK" in val:
                    cell.fill = styles["accent_fill"]
                    cell.font = styles["accent_font"]
                elif val == "GIFTED":
                    cell.fill = styles["done_fill"]
                    cell.font = styles["done_font"]
                elif val == "LOVE":
                    cell.fill = styles["love_fill"]
                    cell.font = styles["love_font"]
                elif val == "LIKE":
                    cell.fill = styles["like_fill"]
                    cell.font = styles["like_font"]
                else:
                    cell.font = styles["data_font"]

    ws4.column_dimensions["A"].width = 28
    ws4.column_dimensions["B"].width = 22
    ws4.column_dimensions["C"].width = 16
    for col_idx in range(4, len(matrix_headers) + 1):
        ws4.column_dimensions[get_column_letter(col_idx)].width = 14

    ws4.freeze_panes = "D2"
    ws4.auto_filter.ref = ws4.dimensions

    wb.save(output_path)
    print(f"[Excel] Successfully exported Gift Planner Report to: {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Daily Bag Loadout Optimizer for Fields of Mistria gift giving, with inventory awareness and Saturday Market support."
    )
    parser.add_argument(
        "--save-file",
        default=None,
        help="Path to .sav file. If omitted, auto-detects the latest save in your FoM saves folder."
    )
    parser.add_argument(
        "--date",
        default=None,
        help="Override simulation date/day (e.g. 'saturday', '6', 'winter 6')"
    )
    parser.add_argument(
        "--mode",
        choices=["auto", "today", "saturday", "market-only", "townsfolk", "weekday", "all"],
        default="auto",
        help="Planning mode: 'auto' (use save day), 'saturday' (simulate Saturday Market), 'market-only' (8 vendors only), 'townsfolk' (26 townsfolk only), 'weekday' (26 townsfolk only), 'all' (all 34 NPCs). Default: auto"
    )
    parser.add_argument(
        "--slots",
        type=int,
        default=20,
        help="Maximum bag slots budgeted for gifts (default: 20)"
    )
    parser.add_argument(
        "--recipes",
        default="data/recipes.json",
        help="Path to recipes.json database (default: data/recipes.json)"
    )
    parser.add_argument(
        "--fiddle",
        default="assets/fiddle",
        help="Path to assets/fiddle directory (default: assets/fiddle)"
    )
    parser.add_argument(
        "--item-data",
        default="data/item_data.json",
        help="Path to item_data.json (default: data/item_data.json)"
    )
    parser.add_argument(
        "--output-dir",
        default="exports",
        help="Output directory for generated CSV and Excel plans (default: exports)"
    )
    parser.add_argument(
        "--format",
        choices=["all", "terminal", "csv", "excel", "both"],
        default="all",
        help="Output format: all, terminal, csv, excel, or both (default: all)"
    )
    parser.add_argument(
        "--loved-weight",
        type=int,
        default=3,
        help="Weight multiplier for loved gifts in optimization scoring (default: 3)"
    )
    parser.add_argument(
        "--liked-weight",
        type=int,
        default=1,
        help="Weight multiplier for liked gifts in optimization scoring (default: 1)"
    )
    parser.add_argument(
        "--vendor-boost",
        type=float,
        default=1.5,
        help="Priority boost multiplier for Saturday Market vendors on market day (default: 1.5)"
    )
    parser.add_argument(
        "--exclude-npcs",
        default="",
        help="Comma-separated list of NPC IDs to exclude from planning"
    )
    parser.add_argument(
        "--force-all-npcs",
        action="store_true",
        help="Force planning for all NPCs, even if save indicates they were already gifted today"
    )
    parser.add_argument(
        "--csv-name",
        default="daily_gift_bag_plan.csv",
        help="Filename for CSV export (default: daily_gift_bag_plan.csv)"
    )
    parser.add_argument(
        "--excel-name",
        default="daily_gift_bag_plan.xlsx",
        help="Filename for Excel export (default: daily_gift_bag_plan.xlsx)"
    )

    args = parser.parse_args()

    repo_root = Path(__file__).resolve().parent
    if not (repo_root / "data").exists() and (repo_root.parent / "data").exists():
        repo_root = repo_root.parent
    fiddle_path = (repo_root / args.fiddle).resolve() if not Path(args.fiddle).is_absolute() else Path(args.fiddle)
    item_data_path = (repo_root / args.item_data).resolve() if not Path(args.item_data).is_absolute() else Path(args.item_data)
    output_dir = (repo_root / args.output_dir).resolve() if not Path(args.output_dir).is_absolute() else Path(args.output_dir)
    recipes_path = (repo_root / args.recipes).resolve() if not Path(args.recipes).is_absolute() else Path(args.recipes)

    # 1. Resolve save file
    save_path = None
    if args.save_file:
        p = Path(args.save_file)
        if p.exists():
            save_path = p.resolve()
        elif (repo_root / args.save_file).exists():
            save_path = (repo_root / args.save_file).resolve()
        else:
            print(f"Error: Specified save file '{args.save_file}' not found.", file=sys.stderr)
            sys.exit(1)
    else:
        # Auto-detect latest save in FoM saves directory or local repo
        latest = find_latest_save()
        local_sample = repo_root / "game-2026225748-autosave.sav"
        if latest and latest.exists():
            save_path = latest
        elif local_sample.exists():
            save_path = local_sample

    save = None
    if save_path and save_path.exists():
        print(f"Loading save file: {save_path}...")
        try:
            save = parse_save_file(save_path)
        except Exception as e:
            print(f"Warning: Failed to parse save file '{save_path}': {e}", file=sys.stderr)
            save = None
    else:
        print("Note: No save file found or specified. Planning with empty starting inventory.")

    # 2. Load recipes database
    recipes = load_recipes(str(recipes_path))

    # 3. Load gift definitions & metadata
    npcs_def = None
    if fiddle_path.exists():
        print(f"Loading NPC preferences from: {fiddle_path}...")
        npcs_def, _ = load_npc_preferences_from_fiddle(fiddle_path)
    if not npcs_def and item_data_path.exists():
        print("Falling back to item_data.json...")
        npcs_def, _ = load_npc_preferences_from_json(item_data_path)

    if not npcs_def:
        print("Error: Failed to load NPC gift definitions.", file=sys.stderr)
        sys.exit(1)

    tracker_files_path = repo_root / "Fields of Mistria Progress Tracker_files"
    if not tracker_files_path.exists():
        tracker_files_path = None
    metadata = load_item_metadata(fiddle_path, item_data_path, tracker_files_path)

    # 4. Handle date override
    effective_mode = args.mode
    if args.date:
        d_lower = args.date.strip().lower()
        if "sat" in d_lower or d_lower in ("6", "13", "20", "27"):
            if effective_mode == "auto":
                effective_mode = "saturday"
        elif "animal" in d_lower or "winter 10" in d_lower or "10 winter" in d_lower or d_lower in ("winter_10", "animal_festival"):
            yr = 1
            if save is not None and hasattr(save, "in_game_date") and save.in_game_date:
                yr = save.in_game_date.year
            elif save is None:
                cal_days = (yr - 1) * 112 + 3 * 28 + 9
                dummy_entries = {
                    "header": json.dumps({"name": "Player", "calendar_time": cal_days * 86400, "clock_time": 36000}),
                    "player": json.dumps({"name": "Player", "inventory": []}),
                    "npcs": json.dumps({}),
                }
                save = SaveData(Path("simulated.sav"), dummy_entries)
            save.in_game_date = InGameDate(year=yr, season="winter", day=10, time_str="10:00")

    # 5. Exclude NPCs if requested
    excluded = set([x.strip().lower() for x in args.exclude_npcs.split(",") if x.strip()])

    # 6. Plan daily gift bag
    plan_results = plan_daily_gift_bag(
        save=save,
        npc_gift_definitions=npcs_def,
        item_metadata=metadata,
        mode=effective_mode,
        max_slots=args.slots,
        loved_weight=args.loved_weight,
        liked_weight=args.liked_weight,
        vendor_boost=args.vendor_boost,
        exclude_npcs=excluded,
        force_all_npcs=args.force_all_npcs,
        recipes=recipes,
    )

    # 7. Output
    if args.format in ("all", "terminal"):
        print_terminal_plan(save, plan_results)

    if args.format in ("all", "csv", "both"):
        export_plan_to_csv(plan_results, output_dir, args.csv_name)

    if args.format in ("all", "excel", "both"):
        excel_path = output_dir / args.excel_name
        export_plan_to_excel(save, plan_results, npcs_def, metadata, excel_path)


if __name__ == "__main__":
    main()
